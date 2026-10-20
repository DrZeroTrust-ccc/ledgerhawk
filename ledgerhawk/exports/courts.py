"""Court records workbook for a subject screen: the screen's vendors with federal court records and DOJ press releases,
government cases first. Built from the screen's outside-context lookup (see pipeline/courtrecords.py for the grading)."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..pipeline import courtindex
from ..pipeline import courtrecords as cr
from .subjects import PRIVILEGED
from .voi import MONEY, NAVY

FOOTER = "Court records found by name for follow-up, not findings of liability. Confirm each party on the docket."
GRADE_FILL = {cr.GOVERNMENT: "F6E3E6", cr.PARTY: "FBF0DC", cr.MENTION: "F2F2F2"}
HEADERS = ["#", "Grade", "Vendor", "UEI", "Dollars on screen ($)", "Case type", "Vendor's side", "Case or release",
           "Court / office", "Docket no.", "Filed", "Closed", "Nature of suit / cause", "Source", "Link", "Match",
           "Analyst verdict", "Why it matters to the agency"]
WIDTHS = [5, 22, 34, 15, 16, 26, 14, 60, 28, 18, 12, 12, 34, 22, 50, 26, 14, 70]
SUMMARY_HEADERS = ["Vendor", "UEI", "Dollars on screen ($)", "Government cases", "Party cases", "Name only",
                   "Latest record", "Top record", "Court search"]
SUMMARY_WIDTHS = [34, 15, 16, 12, 12, 12, 12, 70, 50]


def _header(ws, headers: list[str], widths: list[int], row: int = 1) -> None:
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = f"A{row + 1}"


def _match(r: dict) -> str:
    conf = r.get("confidence")
    return f"{cr.cx.CONFIDENCE[conf]} {r.get('score', 0)}/100" if conf in cr.cx.CONFIDENCE else ""


def build_court_records(screen: dict, generated_at: datetime | None = None, *, include_mentions: bool = False) -> bytes:
    generated_at = generated_at or datetime.now(timezone.utc)
    res = cr.screen_records(screen)
    idx = courtindex.status()
    meta = screen.get("meta") or {}
    wb = Workbook()

    ws = wb.active
    ws.title = "Court Records"
    _header(ws, HEADERS, WIDTHS)
    n = 0
    for v in res["vendors"]:
        for r in v["records"]:
            if r["grade"] == cr.MENTION and not include_mentions:
                continue
            case = r.get("case") or {}
            n += 1
            nature = " · ".join(x for x in (case.get("suit"), case.get("cause")) if x)
            row = [n, cr.GRADES[r["grade"]], v["name"], v["uei"], v["tot"] or None, r["case_type"], r.get("role") or "",
                   r.get("title", ""), case.get("court") or r.get("where", ""), case.get("docket", ""),
                   case.get("filed") or r.get("date", ""), case.get("terminated", ""), nature, r.get("source", ""),
                   r.get("url", ""), _match(r), {"same": "Same entity", "unsure": "Unsure"}.get(r["analyst_verdict"], ""),
                   r["why_it_matters"]]
            ws.append(row)
            i = ws.max_row
            for col in range(1, len(HEADERS) + 1):
                cell = ws.cell(row=i, column=col)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if col == 2:
                    cell.fill = PatternFill("solid", fgColor=GRADE_FILL[r["grade"]])
            ws.cell(row=i, column=5).number_format = MONEY
            if r.get("url"):
                ws.cell(row=i, column=15).hyperlink = r["url"]
    if not n:
        ws.append(["", "No government or party court records found for this screen's vendors."])

    sm = wb.create_sheet("By Vendor")
    _header(sm, SUMMARY_HEADERS, SUMMARY_WIDTHS)
    for v in res["vendors"]:
        top = next((r for r in v["records"] if r["grade"] != cr.MENTION), None)
        latest = max((r.get("date") or "" for r in v["records"] if r["grade"] != cr.MENTION), default="")
        sm.append([v["name"], v["uei"], v["tot"] or None, v["counts"][cr.GOVERNMENT], v["counts"][cr.PARTY],
                   v["counts"][cr.MENTION], latest, f"{top['case_type']}: {top['title']}" if top else "",
                   v["court_error"] or f"Checked {v['fetched_at'][:10]}"])
        sm.cell(row=sm.max_row, column=3).number_format = MONEY

    rm = wb.create_sheet("Read Me")
    rm.column_dimensions["A"].width = 28
    rm.column_dimensions["B"].width = 110
    c = res["counts"]
    lines = [
        ("LedgerHawk court records", ""),
        (PRIVILEGED, ""),
        ("Screen", " · ".join(x for x in (meta.get("matter"), meta.get("client"), meta.get("id")) if x)),
        ("Generated", generated_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Vendors looked up", str(c["vendors"])),
        ("With a government case", str(c[cr.GOVERNMENT])),
        ("With a case as a party", str(c[cr.PARTY])),
        ("Court search incomplete", f"{c['unchecked']} vendors: the federal court search didn't answer (usually its rate "
                                    "limit). Use Re-check courts on the screen to fill them in; they are listed on the "
                                    "By Vendor sheet."),
        ("", ""),
    ] + [(cr.GRADES[g], cr.GRADE_MEANING[g]) for g in cr.GRADE_ORDER] + [
        ("", ""),
        ("Sources", "Federal dockets and opinions from CourtListener (RECAP), and DOJ press releases. State courts, "
                    "PACER-only dockets, GAO bid protests and board of contract appeals decisions not on CourtListener "
                    "are not covered; the subject screen's manual links open those searches."),
        ("Court index", courtindex.summary_line(idx) or "Not loaded on this server: court searches went to "
                        "CourtListener's search API, which allows about 125 searches a day."),
        ("Matching", "Court records name parties, not UEIs. A row is graded on whether the vendor's legal name (or a "
                     "known other name) appears as a party or in the caption. Confirm the party's address or officers "
                     "on the docket before relying on a row; the Match column scores the visible detail."),
        ("Left out", "Records an analyst ruled 'not our subject', records from muted sites"
                     + ("" if include_mentions else ", and records where the name appears but not as a party") + "."),
        ("", ""),
        (FOOTER, ""),
    ]
    for a, b in lines:
        rm.append([a, b])
        rm.cell(row=rm.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    rm["A1"].font = Font(bold=True, size=14, color=NAVY)
    if c["unchecked"]:  # a partial search reads as "clean" unless it says so loudly
        for row in rm.iter_rows():
            if row[0].value == "Court search incomplete":
                row[0].font = row[1].font = Font(bold=True, color="A61B29")
    wb.move_sheet("Read Me", offset=-2)

    wb.properties.creator = "LedgerHawk"
    wb.properties.lastModifiedBy = "LedgerHawk"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
