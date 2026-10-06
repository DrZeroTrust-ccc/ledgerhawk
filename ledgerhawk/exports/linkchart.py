"""Link chart export for a subject screen: an entity list and a link table that i2 Analyst's Notebook and Maltego import.

Entities are organizations (by UEI), people (SAM contacts and screened people), registered suites and exclusion records.
Links say how two entities are connected and where that came from. The data sheets start with their header row so they
import without editing; the privilege label, sources and import steps are on the Read Me sheet.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..pipeline.sam import person_key
from ..pipeline.subjects import POC_ROLE_TEXT
from .subjects import PRIVILEGED
from .voi import FOOTER, NAVY

ENTITY_HEADERS = ["Entity ID", "Entity Type", "Label", "UEI", "Description", "Place", "Excluded", "Subject #", "Notes"]
ENTITY_WIDTHS = [34, 14, 40, 15, 40, 36, 10, 10, 50]
LINK_HEADERS = ["From ID", "From Type", "From Label", "Link Type", "To ID", "To Type", "To Label", "Detail"]
LINK_WIDTHS = [34, 14, 36, 32, 34, 14, 36, 60]
EXCL_LINK = {
    "direct": "Excluded under this UEI",
    "alias": "Named as an alias in exclusion record",
    "jv_partner": "Joint venture carrying the excluded party's name",
    "name_match": "Same name as excluded party (unconfirmed)",
}


class _Chart:
    def __init__(self) -> None:
        self.entities: dict[str, dict] = {}
        self.links: list[dict] = []

    def entity(self, eid: str, kind: str, label: str, **attrs) -> str:
        e = self.entities.setdefault(eid, {"id": eid, "type": kind, "label": label, "uei": "", "description": "",
                                           "place": "", "excluded": False, "subjects": set(), "notes": ""})
        for k, v in attrs.items():
            if k == "subject" and v:
                e["subjects"].add(v)
            elif k == "excluded":
                e["excluded"] = e["excluded"] or bool(v)
            elif v and not e.get(k):
                e[k] = v
        return eid

    def link(self, a: str, b: str, kind: str, detail: str = "") -> None:
        if not any(x["from"] == a and x["to"] == b and x["type"] == kind for x in self.links):
            self.links.append({"from": a, "to": b, "type": kind, "detail": detail})


def _org(ch: _Chart, uei: str, name: str, **attrs) -> str:
    return ch.entity(f"ORG:{uei}", "Organization", name or uei, uei=uei, **attrs)


def _excl(ch: _Chart, h: dict) -> str:
    eid = f"EXCL:{h.get('uei') or h['name']}|{h['agency']}|{h['active_date']}"
    return ch.entity(eid, "Exclusion", f"{h['name']} ({h['agency']} exclusion)", uei=h.get("uei", ""),
                     description=f"{h['type']}, since {h['active_date']}"
                     + (f", until {h['termination_date']}" if h.get("termination_date") else ""),
                     place=", ".join(x for x in [h.get("city", ""), h.get("state", "")] if x), excluded=True,
                     notes=h.get("comments", ""))


def build_chart(screen: dict) -> _Chart:
    ch = _Chart()
    for s in screen["subjects"]:
        ref = s["ref"]
        subject_ids = []
        for e in s["entities"]:
            card = e.get("sam") or {}
            excluded = "EXCLUDED" in e["exclusion_flags"]
            if e["uei"]:
                oid = _org(ch, e["uei"], card.get("legal_name") or e["name"], subject=ref, excluded=excluded,
                           description=f"Subject {ref}: {s['status_label']}" + (f" ({s['role']})" if s["role"] else ""),
                           place=card.get("address", ""))
            else:
                oid = ch.entity(f"ORG:REF{ref}", "Organization", e["name"], subject=ref,
                                description=f"Subject {ref}: {s['status_label']} (no SAM registration found)")
            subject_ids.append(oid)
            persons = {}
            for p in card.get("pocs", []):
                hub = p["universe"] > 5
                pid = ch.entity(f"PER:{p['pkey']}", "Person", p["name"], place=", ".join(x for x in [p["city"].title(), p["state"]] if x),
                                description=p.get("title", ""),
                                notes=f"Contact on {p['universe']} SAM registrations; likely a registered agent or shared office" if hub else "")
                persons[p["pkey"]] = pid
                ch.link(oid, pid, "Has SAM contact", POC_ROLE_TEXT.get(p["role"], p["role"]))
            suite = None
            if card.get("akey"):
                suite = ch.entity(f"ADDR:{card['akey']}", "Address", card["address"],
                                  notes=f"{card['suite_count']} SAM registrations at this suite" if card.get("suite_count", 0) > 1 else "")
                ch.link(oid, suite, "Registered at")
            for h in e["exclusion"]:
                xid = _excl(ch, h)
                if h["kind"] in EXCL_LINK:
                    ch.link(oid, xid, EXCL_LINK[h["kind"]], h.get("support", "") if h["kind"] == "name_match" else "")
                elif h["kind"] == "address":
                    ch.link(xid, suite or oid, "Excluded party at same suite", h.get("evidence", ""))
                elif h["kind"] == "person":
                    ch.link(persons.get(h.get("pkey", ""), oid), xid, "Contact appears on exclusion record", h.get("evidence", ""))
        for r in s["related"]:
            rid = _org(ch, r["uei"], r["name"], excluded=r["excluded"], description="Related entity")
            keys = r.get("via_keys")
            for su in r["of"]:
                sid = f"ORG:{su}"
                if keys is None:  # screens saved before via_keys was recorded
                    ch.link(sid, rid, "Related entity", "; ".join(r["via"]))
                    continue
                for k in keys:
                    if k.startswith("person:") and f"PER:{k[7:]}" in ch.entities:
                        ch.link(rid, f"PER:{k[7:]}", "Has SAM contact", "shared with a subject")
                    elif k == "suite":
                        suite = next((x["to"] for x in ch.links if x["from"] == sid and x["type"] == "Registered at"), None)
                        if suite:
                            ch.link(rid, suite, "Registered at", "same suite as a subject")
                    elif k == "name":
                        ch.link(sid, rid, "Same legal name, another UEI")
            for h in r["exclusion"]:
                ch.link(rid, _excl(ch, h), EXCL_LINK.get(h["kind"], "Exclusion record"))

    for x in screen.get("people") or []:
        key = person_key(x["first"], x["last"], x["state"]) if x["state"] else ""
        pid = ch.entity(f"PER:{key}" if key else f"PER:SCREEN{x['ref']}", "Person", x["input"].split(",")[0].strip(),
                        place=x["state"], description=f"Screened person {x['ref']}: {x['status_label']}",
                        notes="" if x["state"] else "Searched without a state; matches may be different people")
        for r in x["registrations"]:
            ch.link(_org(ch, r["uei"], r["name"], excluded=r["excluded"], place=r["place"]), pid, "Has SAM contact",
                    ", ".join(r["roles"]))
        for h in x["exclusions"]:
            xid = _excl(ch, h)
            ch.link(pid, xid, "Same name as excluded individual" + (" (same state)" if h["support"] == "same state" else " (unconfirmed)"),
                    h["support"])
    return ch


def _sheet(ws, headers: list[str], widths: list[int], rows: list[list]) -> None:
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for r, vals in enumerate(rows, start=2):
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c, v).alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows), 1) + 1}"


def build_linkchart(screen: dict, generated_at: datetime | None = None) -> bytes:
    generated_at = generated_at or datetime.now(timezone.utc)
    m = screen["meta"]
    src = screen["sources"]
    ch = build_chart(screen)
    ents = ch.entities

    wb = Workbook()
    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 120
    lines = []
    if m.get("privileged"):
        lines.append((PRIVILEGED, Font(bold=True, color="A3243B")))
    lines += [
        ("LedgerHawk link chart" + (" (SYNTHETIC DATA)" if m.get("data_class") == "synthetic" else ""), Font(bold=True, size=14, color=NAVY)),
        (f"Matter: {m.get('matter') or 'not set'}" + (f" ({m['client']})" if m.get("client") else "") + f". Screen {m['id']}.", None),
        ("Sources: " + "; ".join(x for x in [
            f"SAM.gov entity extract {src['sam_file']} as of {src['sam_extract_date']}" if src.get("sam_file") else "",
            f"SAM.gov exclusions extract {src['exclusions_file']} as of {src['exclusions_extract_date']}" if src.get("exclusions_file") else "",
        ] if x) + f". Exported {generated_at:%Y-%m-%d %H:%M} UTC.", None),
        (f"{len(ents)} entities and {len(ch.links)} links.", None),
        ("", None),
        ("Entities: one row per organization, person, registered suite or exclusion record. Entity ID is stable across "
         "exports of the same data (ORG:<UEI>, PER:<first|last|state>, ADDR:<normalized suite>, EXCL:<party|agency|date>).", None),
        ("Links: one row per connection, with both ends' IDs, types and labels. This is the sheet to import.", None),
        ("i2 Analyst's Notebook: Import, choose this workbook and the Links sheet, map From ID / From Label / From Type to the "
         "first entity, To ID / To Label / To Type to the second, and Link Type to the link label. Use Entity ID as the identity "
         "so repeated entities merge.", None),
        ("Maltego: Import Graph from Table, choose the Links sheet, and map From Label and To Label to entities (Organization, "
         "Person, Location or Phrase) with Link Type as the link label.", None),
        ("Contacts shared by more than 5 SAM registrations are kept with a note but are likely registered agents or shared "
         "offices. A link is a lead to test, not proof of common ownership or control.", None),
        (FOOTER, None),
    ]
    for n, (text, font) in enumerate(lines, start=1):
        c = ws.cell(n, 1, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            c.font = font

    ws = wb.create_sheet("Entities")
    _sheet(ws, ENTITY_HEADERS, ENTITY_WIDTHS, [
        [e["id"], e["type"], e["label"], e["uei"], e["description"], e["place"], "Yes" if e["excluded"] else "",
         ", ".join(str(x) for x in sorted(e["subjects"])), e["notes"]]
        for e in sorted(ents.values(), key=lambda e: (["Organization", "Person", "Address", "Exclusion"].index(e["type"]), e["label"]))
    ])
    ws = wb.create_sheet("Links")
    _sheet(ws, LINK_HEADERS, LINK_WIDTHS, [
        [x["from"], ents[x["from"]]["type"], ents[x["from"]]["label"], x["type"], x["to"], ents[x["to"]]["type"],
         ents[x["to"]]["label"], x["detail"]] for x in ch.links
    ])

    p = wb.properties
    p.creator = p.lastModifiedBy = "LedgerHawk"
    p.title = "LedgerHawk Link Chart"
    p.created = p.modified = generated_at.replace(tzinfo=None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
