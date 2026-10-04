"""Small-Vendor Screen workbook (Stage 9 integrity lane): Summary, Small-Vendor Leads, Excluded Small Vendors,
Checked and Cleared, Read Me. Metadata names LedgerHawk only."""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..pipeline.integrity import INTEGRITY_MEANING, INTEGRITY_TIERS
from .voi import FOOTER, MONEY, MUTED, NAVY, _disposition, _place

TIER_FILL = {"A": "F6E3E6", "B": "EFE6F4", "C": "FBF0DC", "D": "E4EAF2"}
LEAD_HEADERS = ["#", "Tier", "Vendor UEI", "Vendor Name", "City, State", "FY24 Obligated ($)", "FY25 Obligated ($)",
                "After Exclusion ($)", "Why It Is on This List", "Second Signals", "Excluding Agency", "Routes To", "Analyst Disposition"]
LEAD_WIDTHS = [5, 30, 15, 32, 18, 15, 15, 15, 70, 45, 16, 30, 28]
EXCL_HEADERS = ["#", "Vendor UEI", "Vendor Name", "Excluding Agency", "Excluded Since", "FY24 Obligated ($)", "FY25 Obligated ($)",
                "After Exclusion ($)", "Tier", "Note"]
EXCL_WIDTHS = [5, 15, 32, 16, 14, 15, 15, 15, 30, 70]


def _head(ws, row: int, headers: list[str], widths: list[int]) -> None:
    ws.column_dimensions["A"].width = 3
    for c, (h, w) in enumerate(zip(headers, widths), start=2):
        ws.column_dimensions[get_column_letter(c)].width = w
        cell = ws.cell(row, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _title(ws, text: str, sub: str) -> None:
    ws["B2"] = text
    ws["B2"].font = Font(bold=True, size=16, color=NAVY)
    ws["B3"] = sub
    ws["B3"].font = Font(size=10, color=MUTED)


def _rows(ws, start: int, rows: list[list], money_cols: set[int], height: int = 60) -> int:
    thin = Side(style="thin", color="D9DDE3")
    wrap = Alignment(wrap_text=True, vertical="top")
    for n, vals in enumerate(rows):
        r = start + n
        for c, val in enumerate(vals, start=2):
            cell = ws.cell(r, c, val)
            cell.alignment = wrap
            cell.border = Border(bottom=thin)
            if c in money_cols:
                cell.number_format = MONEY
        ws.row_dimensions[r].height = height
    return start + max(len(rows), 1) - 1


def build_small(items: list[dict], summary: dict, run: dict, generated_at: datetime | None = None) -> bytes:
    """items: [{"v", "wf", "disposition"}] for every vendor with an integrity record; summary: integrity_summary()."""
    generated_at = generated_at or datetime.now(timezone.utc)
    meta = run.get("meta", {})
    man = run.get("manifest", {})
    synthetic = meta.get("data_class") == "synthetic"
    sub = (f"{FOOTER} Run {meta.get('id', '')}, exported {generated_at:%Y-%m-%d %H:%M} UTC. Fiscal-year timing only: "
           "award dates have not been checked against USAspending.")
    order = {"A": 0, "B": 1, "C": 2, "D": 3, "": 4}
    items = sorted(items, key=lambda i: (order[i["v"]["integrity"]["tier"]], -i["v"]["tot"]))
    cleared = [i for i in items if (i.get("disposition") or {}).get("value", "").startswith("Clear")]
    leads = [i for i in items if i["v"]["integrity"]["tier"] and i not in cleared]
    excluded = [i for i in items if i["v"]["integrity"]["excluded"]]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _title(ws, "LedgerHawk Small-Vendor Screen" + (" (SYNTHETIC DATA)" if synthetic else ""), sub)
    _head(ws, 5, ["Step", "Vendors", "Dollars"], [60, 12, 18])
    r = _rows(ws, 6, [[f["label"], f["vendors"], f["dollars"]] for f in summary["funnel"]], {4}, 20)
    _head(ws, r + 2, ["Tier", "Vendors", "Dollars", "After Exclusion", "What It Means"], [60, 12, 18, 18, 80])
    r = _rows(ws, r + 3, [[t["label"], t["vendors"], t["dollars"], t["after_exclusion"], t["meaning"]] for t in summary["tiers"]], {4, 5}, 30)
    ws.cell(r + 2, 2, "Control gaps by excluding agency").font = Font(bold=True, size=12, color=NAVY)
    _head(ws, r + 3, ["Agency", "Tier A", "Tier B", "Tier C", "Excluded", "After Exclusion", "What It Suggests"], [60, 12, 18, 18, 12, 16, 90])
    _rows(ws, r + 4, [[g["agency"], g["A"], g["B"], g["C"], g["excluded"], g["after_exclusion"], g["summary"]] for g in summary["gaps"]], {7}, 45)

    ws = wb.create_sheet("Small-Vendor Leads")
    _title(ws, "Small-Vendor Leads", sub)
    _head(ws, 5, LEAD_HEADERS, LEAD_WIDTHS)
    body = []
    for n, i in enumerate(leads, start=1):
        v, integ = i["v"], i["v"]["integrity"]
        body.append([n, INTEGRITY_TIERS[integ["tier"]], v["uei"], v["name"], _place(v), v["fy24"], v["fy25"], integ["after_exclusion"],
                     "; ".join(integ["reasons"]), "; ".join(integ["second"]), integ["agency"], i["wf"]["owner"], _disposition(i.get("disposition"))])
    last = _rows(ws, 6, body, {7, 8, 9}, 60)
    for n, i in enumerate(leads):
        ws.cell(6 + n, 3).fill = PatternFill("solid", fgColor=TIER_FILL[i["v"]["integrity"]["tier"]])
        ws.cell(6 + n, 3).font = Font(bold=True, color=NAVY)
    ws.freeze_panes = "F6"
    ws.auto_filter.ref = f"B5:N{last}"

    ws = wb.create_sheet("Excluded Small Vendors")
    _title(ws, "Excluded Small Vendors", "Every excluded vendor in the small-vendor lane, paid after exclusion or not.")
    _head(ws, 5, EXCL_HEADERS, EXCL_WIDTHS)
    body = [[n, i["v"]["uei"], i["v"]["name"], i["v"]["integrity"]["agency"], i["v"]["integrity"]["excluded_on"], i["v"]["fy24"],
             i["v"]["fy25"], i["v"]["integrity"]["after_exclusion"],
             INTEGRITY_TIERS.get(i["v"]["integrity"]["tier"], "Not tiered"), "; ".join(i["v"]["integrity"]["reasons"])]
            for n, i in enumerate(excluded, start=1)]
    last = _rows(ws, 6, body, {7, 8, 9}, 45)
    ws.auto_filter.ref = f"B5:K{last}"

    ws = wb.create_sheet("Checked and Cleared")
    _title(ws, "Checked and Cleared", "Small vendors an analyst cleared, with the note that explains why.")
    _head(ws, 5, ["#", "Vendor UEI", "Vendor Name", "Tier Before Review", "Disposition", "Why It Was on the List"], [5, 15, 32, 30, 60, 70])
    _rows(ws, 6, [[n, i["v"]["uei"], i["v"]["name"], INTEGRITY_TIERS.get(i["v"]["integrity"]["tier"], "Not tiered"),
                   _disposition(i.get("disposition")), "; ".join(i["v"]["integrity"]["reasons"])] for n, i in enumerate(cleared, start=1)], set(), 45)

    ws = wb.create_sheet("Read Me")
    _title(ws, "How This Screen Works", FOOTER)
    ws.column_dimensions["B"].width = 120
    notes = [
        "Population: vendors under the materiality line (under $250K combined FY24 and FY25), with closeout artifacts left out. "
        "Growth and concentration signals do not work at this size, so this lane uses integrity signals only.",
        *[f"Tier {k}: {INTEGRITY_MEANING[k]}." for k in INTEGRITY_TIERS],
        "Tier A uses fiscal-year timing: obligations in a fiscal year that began after the exclusion date. FY-level dollars can include "
        "modifications and deobligations, so confirm each Tier A vendor against award dates before acting.",
        "Weak signals never tier a vendor on their own: shared addresses, contacts on many registrations, several UEIs under one name, "
        "and a SAM start date after the first obligations. They appear as second signals.",
        "Not covered: splitting purchases to stay under the micro-purchase threshold. That needs transaction-level data.",
        "What this lane is for: the value is in control gaps, not dollars. Excluded vendors still being paid and excluded parties "
        "with a second door open are questions for the excluding agency's suspension and debarment official and the buying office.",
        "Sources: vendor file " + str(man.get("input_file", "")) + (f"; SAM.gov exclusions extract ({man['exclusions_extract_date']})"
                                                               if man.get("exclusions_extract_date") else "")
        + (f"; SAM.gov entity extract ({man['sam_extract_date']})" if man.get("sam_extract_date") else "") + ".",
        FOOTER,
    ]
    for n, text in enumerate(notes):
        c = ws.cell(5 + n, 2, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[5 + n].height = 32

    p = wb.properties
    p.creator = p.lastModifiedBy = "LedgerHawk"
    p.title = "LedgerHawk Small-Vendor Screen"
    p.created = p.modified = generated_at.replace(tzinfo=None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
