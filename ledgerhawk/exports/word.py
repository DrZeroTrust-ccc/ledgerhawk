"""Editable Word (DOCX) reports for investigators: a vendor case file and a subject screen report.

Investigators write their own reports, so these are drafts to edit, not finished deliverables. Each carries the
matter, the sources with their as-of dates, an optional "Privileged and Confidential" header, and the screening
footer on every page. Document properties name LedgerHawk, never the tool or the analyst's machine.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..pipeline.explain import QUEUE_LABELS, why_it_flagged
from ..pipeline.normalize import money
from ..pipeline.subjects import STATUSES, next_steps
from ..pipeline.tiering import TIERS, category
from .case import POC_ROLE, TIE_LABEL
from .subjects import award_line, awards_for, context_items, context_line, context_summary, note_byline, notes_for, signoff_lines
from .voi import FOOTER, carried_note

PRIVILEGED = "Privileged and Confidential. Prepared at the direction of counsel."
NAVY = RGBColor(0x1F, 0x2A, 0x3A)
MUTED = RGBColor(0x5A, 0x62, 0x70)
CRIMSON = RGBColor(0xA3, 0x24, 0x3B)


def _doc(title: str, privileged: bool, generated_at: datetime) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        doc.styles[name].font.color.rgb = NAVY
    sec = doc.sections[0]
    if privileged:
        hp = sec.header.paragraphs[0]
        r = hp.add_run(PRIVILEGED)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = CRIMSON
    fp = sec.footer.paragraphs[0]
    r = fp.add_run(FOOTER)
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    p = doc.core_properties
    p.author = p.last_modified_by = "LedgerHawk"
    p.title = title
    p.subject = FOOTER
    p.comments = ""
    p.keywords = ""
    p.category = ""
    p.revision = 1
    p.created = p.modified = generated_at.replace(tzinfo=None)
    return doc


def _small(doc: Document, text: str) -> None:
    r = doc.add_paragraph().add_run(text)
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED


def _kv(doc: Document, rows: list[tuple[str, object]]) -> None:
    rows = [(k, v) for k, v in rows if v not in (None, "", [])]
    if not rows:
        return
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text = k
        cells[1].text = str(v)
        for run in cells[0].paragraphs[0].runs:
            run.font.color.rgb = MUTED
            run.font.size = Pt(9)
    doc.add_paragraph()


def _grid(doc: Document, head: list[str], body: list[list[object]]) -> None:
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, head):
        c.text = h
        for run in c.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    _repeat_header(t.rows[0])
    for row in body:
        cells = t.add_row().cells
        for c, val in zip(cells, row):
            c.text = "" if val is None else str(val)
            for run in c.paragraphs[0].runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def _repeat_header(row) -> None:
    pr = row._tr.get_or_add_trPr()
    if pr.find(qn("w:tblHeader")) is None:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        pr.append(el)


def _bullets(doc: Document, items: list[str], numbered: bool = False) -> None:
    for x in items:
        doc.add_paragraph(x, style="List Number" if numbered else "List Bullet")


def _notes(doc: Document, notes: list[dict], placeholder: str) -> None:
    if not notes:
        doc.add_paragraph(placeholder)
    for n in notes:
        if n["text"]:
            doc.add_paragraph(n["text"])
        _small(doc, note_byline(n))


def _context(doc: Document, cx: dict | None, level: int = 3) -> None:
    """Outside context: the summary line, then up to 8 reportable items (confirmed first, then strong and possible
    matches). Name-only hits nobody confirmed and hits an analyst ruled out stay out of the report."""
    if not cx:
        return
    doc.add_heading("Outside context", level=level)
    doc.add_paragraph(context_summary(cx))
    _bullets(doc, [context_line(i) for i in context_items(cx, 8, reportable=True)])
    _small(doc, "Each item is labeled with how it was matched. Only items marked Confirmed were checked by an analyst; the "
                "rest are leads to verify before relying on them.")


def _save(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_case_docx(v: dict, wf: dict, disposition: dict | None, history: list[dict], summary: dict, *,
                    matter: str = "", privileged: bool = False, generated_at: datetime | None = None,
                    context: dict | None = None, case: dict | None = None, ledger: dict | None = None) -> bytes:
    """One-vendor case file, the same facts as the PDF, with investigator next steps, the evidence ledger and the analyst's
    notes and sign-off for this run."""
    generated_at = generated_at or datetime.now(timezone.utc)
    meta = summary.get("meta", {})
    man = summary.get("manifest", {})
    synthetic = meta.get("data_class") == "synthetic"
    doc = _doc(f"LedgerHawk case file: {v['name']}", privileged, generated_at)

    _small(doc, ("SYNTHETIC DATA · " if synthetic else "") + "LedgerHawk case file" + (f" · Matter: {matter}" if matter else ""))
    doc.add_heading(v["name"], level=1)
    sub = [f"UEI {v['uei']}", QUEUE_LABELS.get(v.get("queue") or "", ""),
           f"NAICS {v['naics']}" if v.get("naics") else "", f"PSC {v['psc']}" if v.get("psc") else ""]
    _small(doc, " · ".join(s for s in sub if s))

    tier = wf.get("tier") or ""
    tier_text = TIERS.get(tier, "Not tiered")
    if wf.get("tier_change"):
        c = wf["tier_change"]
        tier_text += f" (set by {c['analyst']} on {c['at'][:10]}: {c['reason']})"
    elif tier:
        tier_text += " (pipeline default)"
    disp = (f"{disposition['value']}: {disposition['note']} ({disposition['analyst']}, {disposition['at'][:10]}{carried_note(disposition)})"
            if disposition else "Not yet dispositioned")
    _kv(doc, [
        ("Tier", tier_text), ("Category", category(v, tier)), ("Assigned to", wf.get("assignee") or "Unassigned"),
        ("Disposition", disp),
        ("Dollars under review", f"FY24 {money(v['fy24'])} · FY25 {money(v['fy25'])} · total {money(v['tot'])}"),
    ])

    doc.add_heading("Why it flagged", level=2)
    doc.add_paragraph(why_it_flagged(v))
    if ledger and ledger["rows"]:
        b = ledger["balance"]
        doc.add_heading("Evidence ledger", level=2)
        _small(doc, f"{b['counts']['strengthens']} findings strengthen the lead, {b['counts']['weakens']} weaken it, "
                    f"{b['counts']['context']} are context. Outside items count only once an analyst confirms they are about this vendor.")
        lean = {"strengthens": "Strengthens", "weakens": "Weakens", "context": "Context"}
        _grid(doc, ["Effect", "Finding", "Source"], [[lean[r["lean"]], r["text"], r["source"] + (f" ({r['by']})" if r.get("by") else "")]
                                                    for r in ledger["rows"]])
    doc.add_heading("Next steps", level=2)
    _bullets(doc, next_steps([v], [], bool(man.get("sam_extract_date"))), numbered=True)

    sigs = v.get("signals") or []
    if sigs:
        doc.add_heading("Screening signals", level=2)
        _grid(doc, ["Signal", "Name", "Detail"], [[s["id"], s["label"], s["detail"]] for s in sigs])

    card = v.get("sam")
    if card:
        doc.add_heading("SAM registration", level=2)
        pocs = "; ".join(f"{p.get('name', '')} ({POC_ROLE.get(p.get('role', ''), p.get('role', ''))}, {p.get('city', '')} {p.get('state', '')})"
                         for p in card.get("pocs") or [])
        flags = [x for x, on in (("address looks like an apartment, unit or PO box", card.get("residential")),
                                 ("address looks like a virtual office or mailbox", card.get("virtual"))) if on]
        _kv(doc, [
            ("Status", ("Active" if card.get("active") else "Not active") + f"; registered {card.get('reg_date', '')}, expires {card.get('exp_date', '')}"),
            ("Business start", card.get("start_date")),
            ("Certifications", ", ".join(card.get("certs") or []) or "None"),
            ("Physical address", card.get("address", "") + (f" ({'; '.join(flags)})" if flags else "")),
            ("Entities at this suite", card.get("suite_count")),
            ("Points of contact", pocs),
        ])

    links = v.get("links") or []
    if links:
        doc.add_heading("Linked vendors", level=2)
        _small(doc, "Different companies that share a contact and a suite or building with this vendor. Shared addresses and "
                    "contacts are signals, not proof of common control.")
        _grid(doc, ["Vendor", "Dollars", "Shared"], [[f"{ln.get('name', '')} ({ln.get('uei', '')})", money(ln.get("tot") or 0),
                                                      ln.get("via", "")] for ln in links])

    doc.add_heading("Exclusions", level=2)
    hits = v.get("exclusion") or []
    if hits:
        _grid(doc, ["Tie", "Excluded party", "Agency", "Type", "Evidence"],
              [[TIE_LABEL.get(h.get("kind", ""), h.get("kind", "")), h.get("name") or h.get("uei", ""), h.get("agency", ""),
                f"{h.get('type', '')}; active {h.get('active_date', '')}", h.get("evidence", "") or h.get("comments", "")] for h in hits])
    else:
        doc.add_paragraph("No active exclusion record is tied to this vendor by UEI, name or alias, or by a shared suite or contact.")

    _context(doc, context, level=2)

    if history:
        doc.add_heading("Analyst history", level=2)
        _grid(doc, ["When", "Analyst", "Action"],
              [[h["at"][:16].replace("T", " "), h["analyst"], (f"[Run {h['other_run']['label']}, {h['other_run']['created_at'][:10]}] " if h.get("other_run") else "") + f"{h['action']}: {h['detail']}"] for h in history])

    doc.add_heading("Investigator notes", level=2)
    review = (case or {}).get("review") or {}
    notes = review.get("notes") or []
    if notes:
        _grid(doc, ["When", "Analyst", "Note", "Evidence"],
              [[n["at"][:10], n["analyst"], n.get("text") or "", (f"{n['file']} (SHA-256 {n['file_sha256'][:16]}…)" if n.get("file") else "")
                + (f" Source: {n['source']}" if n.get("source") else "")] for n in notes])
    else:
        doc.add_paragraph("[Add interviews, registry and court searches, and other work here.]")
    if review.get("history"):
        last = review["history"][-1]
        _small(doc, f"Sign-off: {review.get('state_label', '')}, {last['by']} on {last['at'][:10]}"
               + (f" ({last['comment']})" if last.get("comment") else "") + ".")

    src = [f"vendor file {man.get('input_file', '')}"]
    if man.get("sam_extract_date"):
        src.append(f"SAM entity extract {man['sam_extract_date']}")
    if man.get("exclusions_extract_date"):
        src.append(f"SAM exclusions extract {man['exclusions_extract_date']}")
    _small(doc, f"Run {meta.get('id', '')} · rule set {man.get('rule_set_version', '')} · sources: " + ", ".join(src)
           + f" · generated {generated_at:%Y-%m-%d %H:%M} UTC")
    return _save(doc)


def build_subjects_docx(screen: dict, generated_at: datetime | None = None) -> bytes:
    """Subject screen report: scope, summary, then one section per subject with findings and next steps."""
    generated_at = generated_at or datetime.now(timezone.utc)
    m = screen["meta"]
    src = screen["sources"]
    c = screen["counts"]
    synthetic = m.get("data_class") == "synthetic"
    doc = _doc(f"LedgerHawk subject screen: {m.get('matter') or m['id']}", bool(m.get("privileged")), generated_at)

    _small(doc, ("SYNTHETIC DATA · " if synthetic else "") + "LedgerHawk subject screen")
    doc.add_heading(m.get("matter") or "Subject screen", level=1)
    _kv(doc, [
        ("Client", m.get("client")), ("Prepared by", m.get("created_by")), ("Screened", (m.get("created_at") or "")[:10]),
        ("SAM entity extract", f"{src['sam_file']}, as of {src['sam_extract_date']}" if src.get("sam_file") else "Not used"),
        ("SAM exclusions extract", f"{src['exclusions_file']}, as of {src['exclusions_extract_date']}" if src.get("exclusions_file") else "Not used"),
        ("Rule set", src.get("rule_set_version")),
    ])

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(f"{c['subjects']} subjects screened; {c['related']} related SAM registrations found through shared "
                      "contacts, suites or legal names." + (f" {c['people']} people screened." if c.get("people") else ""))
    if c["subjects"]:
        _grid(doc, ["Status", "Subjects"], [[label, c.get(k, 0)] for k, label in STATUSES.items() if c.get(k)])

    doc.add_heading("Review and sign-off", level=2)
    _kv(doc, signoff_lines(screen))
    if notes_for(screen, "screen"):
        doc.add_heading("Notes on the whole screen", level=3)
        _notes(doc, notes_for(screen, "screen"), "")

    ch = screen.get("changes")
    if ch:
        doc.add_heading("What changed since the last check", level=2)
        ps = ch["parent_sources"]
        doc.add_paragraph(
            f"Compared with the screen of {ch['parent_created_at'][:10]}"
            + (f" (SAM {ps['sam_extract_date']}, exclusions {ps['exclusions_extract_date']})" if ps.get("sam_extract_date") and ps.get("exclusions_extract_date") else "")
            + f": {ch['counts']['changed']} subjects changed, {ch['counts']['worse']} got worse, {ch['counts']['unchanged']} unchanged.")
        for r in ch["subjects"] + (ch.get("people") or []):
            p = doc.add_paragraph()
            p.add_run(f"{r['ref']}. {r['name']}: ").bold = True
            p.add_run(r["status_now_label"] if r["direction"] == "same"
                      else f"{r['status_before_label']} → {r['status_now_label']}")
            _bullets(doc, [f"New: {x}" for x in r["added"]] + [f"No longer found: {x}" for x in r["removed"]])

    for s in screen["subjects"]:
        names = " / ".join((e.get("sam") or {}).get("legal_name") or e["name"] for e in s["entities"])
        doc.add_heading(f"{s['ref']}. {names or s['input_name'] or s['input_uei']}", level=2)
        given = " ".join(x for x in [s["input_uei"], s["input_name"]] if x)
        _kv(doc, [("Status", s["status_label"]), ("Given as", given), ("Role", s["role"]), ("How it was matched", s["resolution"]),
                  ("UEI", ", ".join(e["uei"] for e in s["entities"] if e["uei"]))])
        doc.add_heading("What we found", level=3)
        _bullets(doc, s["findings"])
        doc.add_heading("Next steps", level=3)
        _bullets(doc, s["next_steps"], numbered=True)
        if s["related"]:
            doc.add_heading("Related entities", level=3)
            _grid(doc, ["Entity", "UEI", "How linked", "Exclusion"],
                  [[r["name"], r["uei"], "; ".join(r["via"]),
                    "Excluded" if r["excluded"] else ", ".join(r["flags"])] for r in s["related"]])
            if s["related_total"] > len(s["related"]):
                _small(doc, f"{s['related_total'] - len(s['related'])} more in the workbook.")
        found = awards_for(screen, s["ref"])
        if found:
            aw = screen["awards"]
            doc.add_heading("Federal awards (USAspending)", level=3)
            _bullets(doc, [award_line(e) for e in found])
            top = sorted((a | {"uei": e["uei"]} for e in found for a in e["awards"]), key=lambda a: -a["amount"])[:10]
            if top:
                _grid(doc, ["Award ID", "Agency", "Start", "Obligated", "After exclusion"],
                      [[a["award_id"], a["agency"], a["start"], f"${a['amount']:,.0f}", "Yes" if a["after_exclusion"] else ""] for a in top])
            _small(doc, f"USAspending.gov, looked up {aw['fetched_at'][:10]}. Largest awards shown; the workbook lists all retrieved.")
        for cx in [c for c in (screen.get("context") or {}).get("entities", []) if c.get("ref") == s["ref"]]:
            _context(doc, cx)
        doc.add_heading("Investigator notes", level=3)
        _notes(doc, notes_for(screen, f"s:{s['ref']}"), "[Add work done on this subject here.]")

    for x in screen.get("people") or []:
        doc.add_heading(f"Person {x['ref']}. {x['input']}", level=2)
        _kv(doc, [("Status", x["status_label"]), ("Searched as", f"{x['first'].title()} {x['last'].title()}"
                                                  + (f", {x['state']}" if x["state"] else " (any state)"))])
        doc.add_heading("What we found", level=3)
        _bullets(doc, x["findings"])
        doc.add_heading("Next steps", level=3)
        _bullets(doc, x["next_steps"], numbered=True)
        if x["registrations"]:
            doc.add_heading("SAM registrations listing this person", level=3)
            _grid(doc, ["Entity", "UEI", "Role", "Place", "Excluded"],
                  [[r["name"], r["uei"], ", ".join(r["roles"]), r["place"], "Yes" if r["excluded"] else ""] for r in x["registrations"]])
        for cx in [c for c in (screen.get("context") or {}).get("entities", []) if c.get("person_ref") == x["ref"]]:
            _context(doc, cx)
        doc.add_heading("Investigator notes", level=3)
        _notes(doc, notes_for(screen, f"p:{x['ref']}"), "[Add work done on this person here.]")

    doc.add_heading("Method and limits", level=2)
    _bullets(doc, [
        "This is a screen of named subjects against public federal data. It finds leads to test; it does not establish "
        "ownership, control or wrongdoing.",
        "Subjects given by name are matched to every SAM registration with the same normalized legal name (up to 10).",
        "Related entities share a contact person, a suite or a legal name with a subject. Contacts and addresses shared by "
        "more than 5 SAM entities are treated as registered agents or shared offices and not used as links.",
        "Same-name exclusion matches with nothing else in common are shown as unconfirmed.",
        "People are matched to SAM contacts by first and last name, and by state when one is given. SAM lists contacts, "
        "not owners, and names repeat, so confirm identity before attributing a match.",
        ("Award history is from USAspending.gov as of the lookup date shown, largest contracts and IDVs first; it is the "
         "federal reporting record, not the contract file. Not covered: " if screen.get("awards") else
         "Not covered: award-level history, ") + "corporate registries, beneficial ownership, court records and media.",
        "A subject with no hits is not cleared; the sources and dates above are the scope of this check.",
    ])
    _small(doc, f"Screen {m['id']} · generated {generated_at:%Y-%m-%d %H:%M} UTC")
    return _save(doc)
