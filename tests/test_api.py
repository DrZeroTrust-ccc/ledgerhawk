import importlib

import json
import pytest
from fastapi.testclient import TestClient

from ledgerhawk.pipeline.synthetic import make_synthetic


@pytest.fixture(scope="module")
def ctx(tmp_path_factory):
    root = tmp_path_factory.mktemp("app")
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(root / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(root / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    client = TestClient(appmod.app)
    vendors, excl, sam, planted = make_synthetic(root / "in", n=800, seed=3)
    with open(vendors, "rb") as v, open(excl, "rb") as e:
        r = client.post("/api/runs", files={"vendors": v, "exclusions": e},
                        data={"exclusions_date": "2026-10-02", "synthetic": "true", "analyst": "Test Analyst"})
    assert r.status_code == 200, r.text
    yield client, r.json()["id"], planted
    mp.undo()


def test_run_listed_and_tagged(ctx):
    client, run_id, _ = ctx
    runs = client.get("/api/runs").json()
    assert runs[0]["id"] == run_id
    assert runs[0]["data_class"] == "synthetic"


def test_upload_requires_analyst(ctx):
    client, *_ = ctx
    r = client.post("/api/runs", files={"vendors": ("v.csv", b"UEI,Name,FY24,FY25\n")}, data={"analyst": ""})
    assert r.status_code == 400


def test_names_only_list_is_refused_not_screened_as_zero(ctx):
    client, *_ = ctx
    before = len(client.get("/api/runs").json())
    r = client.post("/api/runs", files={"vendors": ("leads.csv", b"Vendor UEI,Vendor Name,Why\nABC123DEF456,ACME LLC,tip\n")},
                    data={"analyst": "Test Analyst"})
    assert r.status_code == 400 and "Subject screens" in r.json()["detail"]
    assert len(client.get("/api/runs").json()) == before


def test_queue_and_vendor(ctx):
    client, run_id, p = ctx
    q = client.get(f"/api/runs/{run_id}/vendors", params={"queue": "any"}).json()
    ueis = {r["uei"] for r in q["rows"]}
    assert {p["s2_strong"], p["priority"], p["excluded_major"]} <= ueis
    v = client.get(f"/api/runs/{run_id}/vendors/{p['excluded_major']}").json()
    assert "active SAM exclusions list" in v["why"]
    assert "overrides" in v["why"]


def test_funnel_drilldown(ctx):
    client, run_id, p = ctx
    rows = client.get(f"/api/runs/{run_id}/vendors", params={"cut_stage": "1d", "limit": 500}).json()["rows"]
    assert p["major_by_name"] in {r["uei"] for r in rows}
    assert all(r["reason_code"] == "MAJOR_AUDITED" for r in rows)


def test_disposition_requires_note(ctx):
    client, run_id, p = ctx
    url = f"/api/runs/{run_id}/vendors/{p['s4']}/disposition"
    assert client.post(url, json={"value": "Review", "note": " ", "analyst": "A"}).status_code == 400
    assert client.post(url, json={"value": "Guilty", "note": "x", "analyst": "A"}).status_code == 400
    r = client.post(url, json={"value": "Review", "note": "Pull the award file", "analyst": "A"})
    assert r.status_code == 200
    v = client.get(f"/api/runs/{run_id}/vendors/{p['s4']}").json()
    assert v["disposition"]["value"] == "Review"
    assert v["history"][0]["action"] == "disposition"


def test_restore_creates_child_run(ctx):
    client, run_id, p = ctx
    r = client.post(f"/api/runs/{run_id}/vendors/{p['closeout']}/restore", json={"note": "Check deob", "analyst": "A"})
    assert r.status_code == 200
    new_id = r.json()["id"]
    s = client.get(f"/api/runs/{new_id}").json()
    assert s["meta"]["parent_id"] == run_id
    assert s["queue_counts"]["restored"] == 1
    v = client.get(f"/api/runs/{new_id}/vendors/{p['closeout']}").json()
    assert v["restored_from"] == "CLOSEOUT_NET"


def test_bad_run_id(ctx):
    client, *_ = ctx
    assert client.get("/api/runs/../../etc").status_code == 404
    assert client.get("/api/runs/nope").status_code == 404


@pytest.fixture(scope="module")
def sam_ctx(ctx, tmp_path_factory):
    client, _, _ = ctx
    root = tmp_path_factory.mktemp("sam_in")
    vendors, excl, sam, planted = make_synthetic(root, n=800, seed=3)
    with open(sam, "rb") as f:
        r = client.post("/api/sources", files={"file": f}, data={"kind": "sam", "as_of": "2026-09-06", "analyst": "Test Analyst"})
    assert r.status_code == 200, r.text
    sam_id = r.json()["id"]
    with open(excl, "rb") as f:
        r = client.post("/api/sources", files={"file": f}, data={"kind": "exclusions", "as_of": "2026-10-02", "analyst": "Test Analyst"})
    ex_id = r.json()["id"]
    with open(vendors, "rb") as v:
        r = client.post("/api/runs", files={"vendors": v},
                        data={"synthetic": "true", "analyst": "Test Analyst", "sam_source": sam_id, "exclusions_source": ex_id})
    assert r.status_code == 200, r.text
    return client, r.json()["id"], planted


def test_sources_listed_with_staleness(sam_ctx):
    client, *_ = sam_ctx
    src = client.get("/api/sources").json()["sources"]
    kinds = {s["kind"]: s for s in src}
    assert kinds["sam"]["as_of"] == "2026-09-06" and kinds["sam"]["stale_after_days"] == 35
    assert kinds["exclusions"]["stale_after_days"] == 2


def test_relationship_queue_and_graph(sam_ctx):
    client, run_id, p = sam_ctx
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["meta"]["sam_date"] == "2026-09-06" and run["queue_counts"]["relationship"] >= 3
    g = client.get(f"/api/runs/{run_id}/vendors/{p['ex_affiliate']}/graph").json()
    kinds = {n["kind"] for n in g["nodes"]}
    assert {"vendor", "person", "suite", "excluded"} <= kinds
    assert len([n for n in g["nodes"] if n["kind"] == "excluded"]) == 3
    assert all(path["hops"] == 2 for path in g["paths_to_excluded"])
    hub = client.get(f"/api/runs/{run_id}/vendors/{p['hub'][0]}/graph").json()
    assert any(n.get("hub") for n in hub["nodes"])
    assert not [n for n in hub["nodes"] if n["kind"] == "vendor" and not n.get("center")]


def test_exclusion_gaps(sam_ctx):
    client, run_id, p = sam_ctx
    gaps = client.get(f"/api/runs/{run_id}/exclusion-gaps").json()
    by_agency = {g["agency"]: {v["uei"] for v in g["vendors"]} for g in gaps}
    assert p["ex_affiliate"] in by_agency["ICE"]
    assert p["ex_person_vendor"] in by_agency["TSA"]
    assert p["name_collision"] not in set().union(*by_agency.values())


def test_default_tiers_and_owners(sam_ctx):
    client, run_id, p = sam_ctx
    v = client.get(f"/api/runs/{run_id}/vendors/{p['ex_affiliate']}").json()
    assert v["tier"] == "3" and v["owner"] == "ICE suspension and debarment official"
    v = client.get(f"/api/runs/{run_id}/vendors/{p['split_nj']}").json()
    assert v["tier"] == "5" and v["owner"] == "SBA 8(a) continuing-eligibility review"
    v = client.get(f"/api/runs/{run_id}/vendors/{p['s2_strong']}").json()
    assert v["tier"] == "" and v["owner"] == "Awarding agency contracting officer"


def test_tier_change_needs_reason_and_is_audited(sam_ctx):
    client, run_id, p = sam_ctx
    url = f"/api/runs/{run_id}/vendors/{p['succ_new']}/tier"
    assert client.post(url, json={"tier": "1", "reason": " ", "analyst": "A"}).status_code == 400
    assert client.post(url, json={"tier": "9", "reason": "x", "analyst": "A"}).status_code == 400
    r = client.post(url, json={"tier": "1", "reason": "Award records show a single IDIQ", "analyst": "Test Analyst"})
    assert r.status_code == 200 and r.json()["prior"] == "5"
    v = client.get(f"/api/runs/{run_id}/vendors/{p['succ_new']}").json()
    assert v["tier"] == "1" and v["tier_default"] == "5" and v["tier_change"]["reason"].startswith("Award records")
    assert v["history"][0]["action"] == "tier" and "Promoted to 1 Elevated (was 5 Not yet reviewed)" in v["history"][0]["detail"]
    rows = client.get(f"/api/runs/{run_id}/vendors", params={"tier": "1"}).json()["rows"]
    assert [r["uei"] for r in rows] == [p["succ_new"]]


def test_routing_assignment_and_rollup(sam_ctx):
    client, run_id, p = sam_ctx
    r = client.post(f"/api/runs/{run_id}/vendors/{p['young']}/routing", json={"owner": "GSA FAS contracting officer", "analyst": "Test Analyst"})
    assert r.status_code == 200
    r = client.post(f"/api/runs/{run_id}/assign", json={"ueis": [p["young"], p["split_nj"]], "assignee": "Pat", "analyst": "Test Analyst"})
    assert r.json() == {"assigned": 2}
    rows = client.get(f"/api/runs/{run_id}/vendors", params={"assignee": "Pat"}).json()["rows"]
    assert {r["uei"] for r in rows} == {p["young"], p["split_nj"]}
    assert next(r for r in rows if r["uei"] == p["young"])["owner"] == "GSA FAS contracting officer"
    roll = client.get(f"/api/runs/{run_id}/tier-rollup").json()
    tiers = {t["tier"]: t for t in roll["tiers"]}
    assert tiers["3"]["vendors"] >= 1 and tiers["5"]["vendors"] >= 2 and roll["assignees"]["Pat"] == 2
    assert sum(roll["dispositions"].values()) >= tiers["5"]["vendors"]


def test_vendors_of_interest_export(sam_ctx):
    import io

    import openpyxl
    client, run_id, p = sam_ctx
    r = client.get(f"/api/runs/{run_id}/exports/vendors-of-interest.xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.properties.creator == "LedgerHawk" and wb.properties.lastModifiedBy == "LedgerHawk"
    ws = wb["Vendors of Interest"]
    assert [c.value for c in ws[5]][1:5] == ["#", "Tier", "Category", "Vendor UEI"]
    assert "not findings of fraud" in ws["B3"].value and "SYNTHETIC" in ws["B2"].value
    rows = [[c.value for c in row][1:] for row in ws.iter_rows(min_row=6) if row[1].value]
    tiers = [r[1] for r in rows]
    assert tiers == sorted(tiers) and all(t[0] in "12345" for t in tiers)
    assert any(r[3] == p["succ_new"] for r in rows)
    text = " ".join(str(c) for r in rows for c in r).lower()
    assert "fraud" not in text and "guilty" not in text
    assert "Read Me" in wb.sheetnames


def test_case_pdf(sam_ctx):
    client, run_id, p = sam_ctx
    r = client.get(f"/api/runs/{run_id}/vendors/{p['succ_new']}/case.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert b"LedgerHawk" in r.content
    assert client.get(f"/api/runs/{run_id}/vendors/NOPE/case.pdf").status_code == 404


def test_access_gate(ctx, monkeypatch):
    import base64

    from ledgerhawk.api import app as appmod
    client = ctx[0]
    monkeypatch.setattr(appmod, "ACCESS_PASSWORD", "s3cret")
    assert client.get("/api/healthz").status_code == 200
    r = client.get("/api/runs")
    assert r.status_code == 401 and "Basic" in r.headers["www-authenticate"]
    bad = base64.b64encode(b"chase:nope").decode()
    assert client.get("/api/runs", headers={"Authorization": f"Basic {bad}"}).status_code == 401
    good = base64.b64encode(b"chase:s3cret").decode()
    assert client.get("/api/runs", headers={"Authorization": f"Basic {good}"}).status_code == 200


def test_seed_synthetic(tmp_path, monkeypatch):
    from ledgerhawk.api import app as appmod
    from ledgerhawk.api.store import Store
    monkeypatch.setattr(appmod, "store", Store(tmp_path))
    appmod._seed_synthetic()
    runs = appmod.store.list_runs()
    assert len(runs) == 1 and runs[0]["data_class"] == "synthetic"
    appmod._seed_synthetic()
    assert len(appmod.store.list_runs()) == 1
    # A run that predates the integrity lane gets a fresh synthetic run beside it.
    rj = appmod.store.run_dir(runs[0]["id"]) / "run.json"
    s = json.loads(rj.read_text())
    s["queue_counts"].pop("integrity_leads")
    rj.write_text(json.dumps(s))
    appmod._seed_synthetic()
    assert len(appmod.store.list_runs()) == 2


def test_integrity_lane_endpoint_and_export(sam_ctx):
    import io

    import openpyxl
    client, run_id, p = sam_ctx
    d = client.get(f"/api/runs/{run_id}/integrity").json()
    tiers = {t["tier"]: t["vendors"] for t in d["tiers"]}
    assert tiers["A"] >= 2 and tiers["B"] >= 1 and tiers["C"] >= 1 and tiers["D"] >= 1
    by = {r["uei"]: r for r in d["rows"]}
    assert by[p["small_paid_after"]]["integrity"]["tier"] == "A"
    assert [r["integrity"]["tier"] for r in d["rows"] if r["integrity"]["tier"]] == sorted(r["integrity"]["tier"] for r in d["rows"] if r["integrity"]["tier"])
    gsa = next(g for g in d["gaps"] if g["agency"] == "GSA")
    assert gsa["A"] >= 1 and "after the exclusion" in gsa["summary"]
    rows = client.get(f"/api/runs/{run_id}/vendors", params={"queue": "integrity"}).json()["rows"]
    assert p["small_suite"] in {r["uei"] for r in rows}
    r = client.get(f"/api/runs/{run_id}/exports/small-vendor-screen.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Small-Vendor Leads", "Excluded Small Vendors", "Checked and Cleared", "Read Me"]
    assert wb.properties.creator == "LedgerHawk"
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value).lower()
    assert "fraud" not in text.replace("not findings of fraud", "") and "guilty" not in text
    v = client.get(f"/api/runs/{run_id}/vendors/{p['small_paid_after']}").json()
    assert "integrity lane, tier A" in v["why"]


def test_case_docx(sam_ctx):
    client, run_id, p = sam_ctx
    r = client.get(f"/api/runs/{run_id}/vendors/{p['succ_new']}/case.docx?matter=M-1&privileged=true")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "wordprocessingml" in r.headers["content-type"]
    assert client.get(f"/api/runs/{run_id}/vendors/NOPE/case.docx").status_code == 404


def test_runs_keep_their_own_decisions_and_follow_ups_carry_them_labeled(sam_ctx):
    client, run_id, p = sam_ctx
    uei = p["s4"]
    # the first run's "Review" on this vendor (test_disposition_requires_note) stays in that run
    assert client.get(f"/api/runs/{run_id}/vendors/{uei}").json()["disposition"] is None
    r = client.post(f"/api/runs/{run_id}/vendors/{uei}/disposition",
                    json={"value": "Clear – lawful explanation", "note": "Competed award", "analyst": "Ana"})
    assert r.status_code == 200
    client.post(f"/api/runs/{run_id}/assign", json={"ueis": [uei], "assignee": "Ben", "analyst": "Ana"})

    assert client.post(f"/api/runs/{run_id}/follow-up", data={"analyst": ""}).status_code == 400
    meta = client.get(f"/api/runs/{run_id}").json()["meta"]

    # a follow-up run carries the decision forward, labeled, until it is confirmed here
    f = client.post(f"/api/runs/{run_id}/follow-up", data={"analyst": "Ana"})
    assert f.status_code == 200, f.text
    fid = f.json()["id"]
    s = client.get(f"/api/runs/{fid}").json()
    assert s["meta"]["follows_id"] == run_id and s["follows"]["label"] == meta["label"]
    assert s["changes"]["counts"] == {"new": 0, "dropped": 0, "changed": 0}
    assert client.get(f"/api/runs/{run_id}").json()["followed_by"][0]["id"] == fid
    v = client.get(f"/api/runs/{fid}/vendors/{uei}").json()
    assert v["disposition"]["carried_from"]["id"] == run_id
    assert v["assignee"] == "Ben"
    rec = client.get(f"/api/runs/{fid}/record").json()
    assert rec["queue"]["carried"] == 1 and rec["queue"]["decided"] == 0
    assert rec["follows"]["id"] == run_id
    assert [h["action"] for h in rec["log"]] == ["run_created"]

    # deciding in the follow-up doesn't touch the earlier run
    client.post(f"/api/runs/{fid}/vendors/{uei}/disposition", json={"value": "Review", "note": "New award", "analyst": "Cy"})
    assert client.get(f"/api/runs/{run_id}/vendors/{uei}").json()["disposition"]["value"] == "Clear – lawful explanation"
    v = client.get(f"/api/runs/{fid}/vendors/{uei}").json()
    assert v["disposition"]["value"] == "Review" and "carried_from" not in v["disposition"]
    earlier = [h for h in v["history"] if h.get("other_run")]
    assert any(h["other_run"]["id"] == run_id for h in earlier)

    # confirming adopts a carried decision in this run and logs where it came from
    uei2 = p["priority"]
    client.post(f"/api/runs/{run_id}/vendors/{uei2}/disposition", json={"value": "Refer", "note": "To OIG", "analyst": "Ana"})
    r = client.post(f"/api/runs/{fid}/confirm-carried", json={"ueis": [uei2, uei], "analyst": "Cy"})
    assert r.json() == {"confirmed": 1}
    d = client.get(f"/api/runs/{fid}/vendors/{uei2}").json()["disposition"]
    assert d["value"] == "Refer" and "carried_from" not in d and d["analyst"] == "Ana"
    log = client.get(f"/api/runs/{fid}/record").json()["log"]
    assert log[0]["action"] == "disposition_confirmed" and "Kept Refer from run" in log[0]["detail"]
    assert all(h["run_id"] == fid for h in log)


def test_follow_up_lists_what_changed():
    from ledgerhawk.api.store import compare_runs
    row = lambda u, q, tot, flags=(), sig=(): {"uei": u, "name": u, "queue": q, "tot": tot, "reason": "Under $250K",  # noqa: E731
                                              "exclusion_flags": list(flags), "signals": [{"id": s} for s in sig]}
    old = [row("A", "signals", 1e6, sig=["S1"]), row("B", "signals", 2e6), row("C", "", 1e5)]
    new = [row("A", "exclusion", 3e6, flags=["EXCLUDED"], sig=["S1", "S4"]), row("B", "", 1e5), row("C", "signals", 4e6),
           row("D", "signals", 5e6)]
    ch = compare_runs(old, new)
    assert ch["counts"] == {"new": 2, "dropped": 1, "changed": 1}
    assert [x["uei"] for x in ch["new"]] == ["D", "C"]
    assert ch["new"][1]["why"].startswith("Earlier set aside")
    assert ch["dropped"][0]["why"] == "Now set aside: Under $250K"
    kinds = {w["kind"] for w in ch["changed"][0]["what"]}
    assert {"flags", "signals", "dollars"} <= kinds


def test_old_decisions_move_to_their_run(tmp_path):
    import sqlite3
    from ledgerhawk.api.store import Store
    db = sqlite3.connect(tmp_path / "state.db")
    db.executescript("""
        CREATE TABLE disposition (uei TEXT PRIMARY KEY, value TEXT NOT NULL, note TEXT NOT NULL, analyst TEXT NOT NULL, at TEXT NOT NULL, run_id TEXT);
        CREATE TABLE routing (uei TEXT PRIMARY KEY, owner TEXT NOT NULL, analyst TEXT NOT NULL, at TEXT NOT NULL);
        CREATE TABLE audit (id INTEGER PRIMARY KEY AUTOINCREMENT, at TEXT NOT NULL, analyst TEXT NOT NULL, action TEXT NOT NULL, uei TEXT, run_id TEXT, detail TEXT);
        INSERT INTO disposition VALUES ('U1', 'Review', 'n', 'A', '2026-10-01', 'R1');
        INSERT INTO routing VALUES ('U1', 'GSA OIG', 'A', '2026-10-01');
        INSERT INTO audit (at, analyst, action, uei, run_id, detail) VALUES ('2026-10-01', 'A', 'routing', 'U1', 'R1', 'x');
    """)
    db.commit()
    db.close()
    s = Store(tmp_path)
    with s._db() as db:
        assert db.execute("SELECT run_id, value FROM run_disposition").fetchall() == [("R1", "Review")]
        assert db.execute("SELECT run_id, owner FROM run_routing").fetchall() == [("R1", "GSA OIG")]
    Store(tmp_path)  # runs once


def test_fast_queue_endpoints(sam_ctx):
    client, run_id, p = sam_ctx
    rows = client.get(f"/api/runs/{run_id}/vendors", params={"queue": "any"}).json()["rows"]
    assert all(r["headline"] for r in rows)
    ex = next(r for r in rows if "EXCLUDED" in r["exclusion_flags"])
    assert ex["headline"].startswith("On the SAM exclusions list")

    ueis = [r["uei"] for r in rows if not r["disposition"]][:3]
    assert client.post(f"/api/runs/{run_id}/dispositions", json={"ueis": ueis, "value": "Review", "note": " ", "analyst": "Q"}).status_code == 400
    assert client.post(f"/api/runs/{run_id}/dispositions", json={"ueis": ["NOPE"], "value": "Review", "note": "x", "analyst": "Q"}).status_code == 404
    before = client.get(f"/api/runs/{run_id}/progress", params={"analyst": "q"}).json()
    r = client.post(f"/api/runs/{run_id}/dispositions", json={"ueis": ueis, "value": "Review", "note": "Batch: pull files", "analyst": "Q"})
    assert r.json() == {"decided": 3}
    after = client.get(f"/api/runs/{run_id}/progress", params={"analyst": "q"}).json()
    assert after["open"] == before["open"] - 3 and after["mine_today"] == before["mine_today"] + 3

    client.post(f"/api/runs/{run_id}/assign", json={"ueis": ueis[:2], "assignee": "Quinn", "analyst": "Q"})
    mine = client.get("/api/my-cases", params={"analyst": "quinn"}).json()["rows"]
    # shown once, in the newest run that continues this one (an earlier test started a follow-up)
    latest = ([r["id"] for r in client.get(f"/api/runs/{run_id}").json()["followed_by"]] or [run_id])[-1]
    assert sorted(m["uei"] for m in mine) == sorted(ueis[:2]) and all(m["run"]["id"] == latest for m in mine)
    assert client.get("/api/my-cases").status_code == 400


def test_case_notes_evidence_signoff_awards_and_ledger(sam_ctx):
    from test_subjects import _fake_usaspending
    import ledgerhawk.api.app as appmod
    client, run_id, p = sam_ctx
    uei = p["excluded_major"]
    base = f"/api/runs/{run_id}/vendors/{uei}"
    v = client.get(base).json()
    assert v["case"]["review"]["state"] == "draft" and v["case"]["awards"] is None
    assert v["ledger"]["rows"][0]["lean"] == "strengthens" and v["ledger"]["balance"]["lean"] == "strengthens"

    assert client.post(f"{base}/notes", data={"analyst": "Ana", "text": "x", "lean": "maybe"}).status_code == 400
    r = client.post(f"{base}/notes", data={"analyst": "Ana", "text": "Owner confirmed same person via state registry",
                                           "lean": "strengthens", "source": "State SOS"},
                    files={"file": ("sos.pdf", b"%PDF-1.4 test")})
    assert r.status_code == 200, r.text
    note = r.json()
    assert note["file_sha256"] and note["lean"] == "strengthens"
    client.post(f"{base}/notes", data={"analyst": "Ana", "text": "Contract was competed with 4 offers", "lean": "weakens"})
    assert client.get(f"{base}/evidence/{note['id']}").content == b"%PDF-1.4 test"
    led = client.get(base).json()["ledger"]
    assert any(r["kind"] == "note" and r["lean"] == "weakens" for r in led["rows"])

    appmod.store.awards_post = _fake_usaspending([])
    try:
        r = client.post(f"{base}/awards", data={"analyst": "Ana"})
    finally:
        appmod.store.awards_post = None
    assert r.status_code == 200 and r.json()["entities"][0]["count"] == 3

    # two-person sign-off locks the notes
    assert client.post(f"{base}/review", data={"analyst": "Ana", "action": "submit"}).json()["state"] == "submitted"
    assert client.post(f"{base}/review", data={"analyst": "Ana", "action": "approve"}).status_code == 400
    assert client.post(f"{base}/review", data={"analyst": "Ben", "action": "approve"}).json()["state"] == "approved"
    assert client.post(f"{base}/notes", data={"analyst": "Ana", "text": "late"}).status_code == 400
    log = client.get(f"/api/runs/{run_id}/record").json()["log"]
    assert {"case_note", "case_awards", "case_submit", "case_approve"} <= {h["action"] for h in log}

    # a follow-up run shows these notes as earlier notes; its own sign-off starts over
    fid = client.post(f"/api/runs/{run_id}/follow-up", data={"analyst": "Ana"}).json()["id"]
    c = client.get(f"/api/runs/{fid}/vendors/{uei}").json()["case"]
    assert c["review"]["state"] == "draft" and not c["review"]["notes"]
    assert len(c["earlier_notes"]) == 2 and c["earlier_notes"][0]["run"]["id"] == run_id
    assert client.get(f"/api/runs/{fid}/vendors/{uei}/evidence/{note['id']}").status_code == 200
    assert client.get(f"{base}/case.docx").status_code == 200
    assert client.get(f"/api/runs/{run_id}/vendors/NOPE/notes").status_code in (404, 405)


class _FakeClaude:
    """Stands in for anthropic.Anthropic(): returns a canned structured answer and records the prompt."""

    def __init__(self, answer: dict, stop_reason: str = "end_turn"):
        import json
        from types import SimpleNamespace as NS
        self.calls = []
        outer = self

        def create(**kw):
            outer.calls.append(kw)
            return NS(stop_reason=stop_reason, model="claude-opus-5-5", content=[NS(type="text", text=json.dumps(answer))])
        self.beta = NS(messages=NS(create=create))


def test_case_summary_drafted_sourced_edited_and_exported(sam_ctx):
    import ledgerhawk.api.app as appmod
    client, run_id, p = sam_ctx
    uei = p["excluded_small"]
    base = f"/api/runs/{run_id}/vendors/{uei}"
    v = client.get(base).json()
    assert v["case"]["summary"] is None and v["ledger"]["rows"][0]["id"] == "E1"

    fake = _FakeClaude({"sentences": [{"text": "Screening ties this vendor to an excluded firm.", "sources": ["E1"]},
                                      {"text": "An unsourced claim.", "sources": []},
                                      {"text": "A claim citing nothing real.", "sources": ["E99"]}],
                        "next_steps": [{"text": "Pull the exclusion record.", "sources": ["E1"]}]})
    appmod.store.summary_client = fake
    try:
        assert client.post(f"{base}/summary/draft", data={"analyst": ""}).status_code == 400
        r = client.post(f"{base}/summary/draft", data={"analyst": "Ana"})
        assert r.status_code == 200, r.text
        s = r.json()
        # only sentences tied to real ledger rows survive
        assert [x["text"] for x in s["sentences"]] == ["Screening ties this vendor to an excluded firm."]
        assert s["requested_by"] == "Ana" and not s["edited_by"] and s["cited"]["E1"]
        sent = fake.calls[0]
        assert "E1 [strengthens]" in sent["messages"][0]["content"] and sent["fallbacks"] == "default"
        got = client.get(base).json()
        assert got["summary_enabled"] and got["case"]["summary"]["stale"] is False

        r = client.post(f"{base}/summary", json={"analyst": "Ben", "sentences": [
            {"text": "Screening links this vendor to an excluded firm with the same owner.", "sources": ["E1"]}, {"text": " ", "sources": []}],
            "next_steps": []})
        assert r.status_code == 200 and r.json()["edited_by"] == "Ben" and len(r.json()["sentences"]) == 1

        # new evidence makes the summary out of date
        client.post(f"{base}/notes", data={"analyst": "Ana", "text": "Registry shows a different owner", "lean": "weakens"})
        assert client.get(base).json()["case"]["summary"]["stale"] is True

        log = {h["action"] for h in client.get(f"/api/runs/{run_id}/record").json()["log"]}
        assert {"case_summary_drafted", "case_summary_edited"} <= log
        import io
        from docx import Document
        text = "\n".join(par.text for par in Document(io.BytesIO(client.get(f"{base}/case.docx").content)).paragraphs)
        assert "same owner. [E1]" in text and "edited by Ben" in text

        appmod.store.summary_client = _FakeClaude({}, stop_reason="refusal")
        assert client.post(f"{base}/summary/draft", data={"analyst": "Ana"}).status_code == 503
    finally:
        appmod.store.summary_client = None


def test_hawk_writes_queue_reasons(sam_ctx):
    import json
    from types import SimpleNamespace as NS
    import ledgerhawk.api.app as appmod
    client, run_id, p = sam_ctx
    calls = []

    def create(**kw):
        calls.append(kw)
        lines = kw["messages"][0]["content"].splitlines()[1:]
        ids = [ln.split(" | ")[0] for ln in lines]
        reasons = [{"id": u, "text": f"Reason for {u}."} for u in ids] + [{"id": "NOT-A-LEAD", "text": "x"}]
        return NS(stop_reason="end_turn", content=[NS(type="text", text=json.dumps({"reasons": reasons}))])

    appmod.store.summary_client = NS(beta=NS(messages=NS(create=create)))
    appmod.store.hawk_inline = True
    try:
        assert client.get(f"/api/runs/{run_id}/hawk-reasons").json()["state"] == "none"
        assert client.post(f"/api/runs/{run_id}/hawk-reasons", data={"analyst": ""}).status_code == 400
        r = client.post(f"/api/runs/{run_id}/hawk-reasons", data={"analyst": "Ana"})
        assert r.status_code == 200, r.text
        st = r.json()
        queued = client.get(f"/api/runs/{run_id}/vendors", params={"queue": "any", "limit": 500}).json()
        assert st["elapsed_s"] >= 0
        assert st["state"] == "done" and st["written"] == queued["total"] == st["total"] > 0
        assert all(row["hawk"] == f"Reason for {row['uei']}." for row in queued["rows"])
        assert "obligations $" in calls[0]["messages"][0]["content"] and calls[0]["output_config"]["effort"] == "low"
        # leads outside the queue get no reason; a second ask only covers leads still missing one
        rest = client.get(f"/api/runs/{run_id}/vendors", params={"limit": 500}).json()["rows"]
        assert all(not row["hawk"] for row in rest if not row["queue"])
        n = len(calls)
        assert client.post(f"/api/runs/{run_id}/hawk-reasons", data={"analyst": "Ana"}).json()["total"] == 0
        assert len(calls) == n
    finally:
        appmod.store.summary_client = None
        appmod.store.hawk_inline = False
