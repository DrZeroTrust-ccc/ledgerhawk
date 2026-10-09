"""Subject screening on named targets, checked against planted synthetic cases."""
import importlib
import io
from datetime import date

import pytest
import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ledgerhawk.exports.subjects import PRIVILEGED, build_subjects
from ledgerhawk.pipeline.exclusions import load_exclusions
from ledgerhawk.pipeline.sam import load_sam
from ledgerhawk.pipeline.subjects import parse_subjects, subject_screen
from ledgerhawk.pipeline.synthetic import make_synthetic

EXCL_DATE = date(2026, 10, 2)
SAM_DATE = date(2026, 9, 6)


@pytest.fixture(scope="module")
def syn(tmp_path_factory):
    d = tmp_path_factory.mktemp("subjects")
    vendors, excl, sam, planted = make_synthetic(d, n=1500, seed=5)
    return load_sam(sam, SAM_DATE), load_exclusions(excl, EXCL_DATE), planted, (vendors, excl, sam)


def by_ref(res):
    return {s["ref"]: s for s in res.subjects}


def test_parse_pasted_and_file(tmp_path):
    subs = parse_subjects("SYNABCDEFGH1, Acme Supply LLC\nAcme Supply LLC\n\nsynabcdefgh1\nBeta Corp")
    assert [(s["uei"], s["name"]) for s in subs] == [("SYNABCDEFGH1", "Acme Supply LLC"), ("", "Acme Supply LLC"), ("", "Beta Corp")]
    f = tmp_path / "list.csv"
    f.write_text("Company Name,Unique Entity ID,Relationship\nGamma LLC,SYNABCDEFGH2,Supplier\n")
    subs = parse_subjects(path=f)
    assert subs[0]["uei"] == "SYNABCDEFGH2" and subs[0]["role"] == "Supplier"
    with pytest.raises(ValueError):
        parse_subjects("")


def test_statuses_and_next_steps(syn):
    sam, ex, p, _ = syn
    text = "\n".join([p["excluded_major"], p["ex_affiliate"], p["alias_target"], p["name_collision"], "ZZZZZZZZZZZZ", p["hub"][0]])
    s = by_ref(subject_screen(parse_subjects(text), sam, ex))
    assert s[1]["status"] == "excluded" and s[1]["next_steps"][0].startswith("Confirm the exclusion")
    assert s[2]["status"] == "tied" and s[2]["related_total"] == 3
    assert all(r["excluded"] for r in s[2]["related"])
    assert s[3]["status"] == "tied"  # named as an alias in an exclusion record
    assert s[4]["status"] == "name_only" and "likely a different firm" in " ".join(s[4]["findings"])
    assert s[5]["status"] == "registration" and s[5]["resolution"] == "UEI not in the SAM extract"
    # A registered-agent address is a hub: no related entities through it.
    assert s[6]["related_total"] == 0 and "registered agent" in " ".join(s[6]["findings"])


def test_name_resolution_and_family(syn):
    sam, ex, p, _ = syn
    s = subject_screen(parse_subjects("Synthetic Debarred Supply LLC\nNo Such Company Anywhere LLC"), sam, ex)
    first, second = by_ref(s)[1], by_ref(s)[2]
    assert first["resolution"].startswith("Name matched 2 SAM registrations") and first["status"] == "excluded"
    assert second["resolution"] == "No SAM registration with this name"
    # A subject given by UEI picks up its other registrations under the same legal name.
    split = by_ref(subject_screen(parse_subjects(p["split_nj"]), sam, ex))[1]
    assert {r["uei"] for r in split["related"]} == {p["split_va"], p["split_md"]}
    assert "$0" not in " ".join(split["findings"])


def test_dollars_from_a_run(syn):
    sam, ex, p, _ = syn
    dollars = {p["succ_new"]: {"fy24": 500_000, "fy25": 6_000_000}}
    s = by_ref(subject_screen(parse_subjects(p["succ_new"]), sam, ex, dollars=dollars))[1]
    assert s["entities"][0]["fy25"] == 6_000_000
    assert any(sig["id"] == "S3" for sig in s["entities"][0]["signals"])


def test_requires_a_source():
    with pytest.raises(ValueError):
        subject_screen(parse_subjects("Acme LLC"))


def test_workbook(syn):
    sam, ex, p, _ = syn
    res = subject_screen(parse_subjects(f"{p['ex_affiliate']}\n{p['excluded_major']}"), sam, ex).to_dict()
    res["meta"] = {"id": "x", "matter": "M-1", "client": "Client", "privileged": True, "created_by": "A", "created_at": "now"}
    wb = load_workbook(io.BytesIO(build_subjects(res)))
    assert wb.sheetnames == ["Summary", "Subjects", "Related Entities", "Exclusion Records", "Read Me"]
    assert all(wb[n]["B1"].value == PRIVILEGED for n in wb.sheetnames)
    assert wb.properties.creator == "LedgerHawk"
    assert wb["Related Entities"].max_row >= 8  # three excluded entities at the shared suite


def test_api(syn, tmp_path):
    _, _, p, (vendors, excl, sam) = syn
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    client = TestClient(appmod.app)
    try:
        sources = {}
        for kind, path, d in (("sam", sam, "2026-09-06"), ("exclusions", excl, "2026-10-02")):
            with open(path, "rb") as f:
                r = client.post("/api/sources", files={"file": f}, data={"kind": kind, "as_of": d, "analyst": "T"})
            assert r.status_code == 200, r.text
            sources[kind] = r.json()["id"]
        data = {"subjects_text": p["ex_affiliate"], "analyst": "T", "matter": "M-1", "privileged": "true",
                "sam_source": sources["sam"], "exclusions_source": sources["exclusions"], "synthetic": "true"}
        assert client.post("/api/subject-screens", data={**data, "analyst": ""}).status_code == 400
        assert client.post("/api/subject-screens", data={**data, "subjects_text": ""}).status_code == 400
        r = client.post("/api/subject-screens", data=data)
        assert r.status_code == 200, r.text
        sid = r.json()["id"]
        listed = client.get("/api/subject-screens").json()
        assert listed[0]["id"] == sid and listed[0]["counts"]["tied"] == 1 and listed[0]["privileged"]
        got = client.get(f"/api/subject-screens/{sid}").json()
        assert got["subjects"][0]["status"] == "tied"
        x = client.get(f"/api/subject-screens/{sid}/subject-screen.xlsx")
        assert x.status_code == 200 and x.content[:2] == b"PK"
        w = client.get(f"/api/subject-screens/{sid}/subject-screen.docx")
        assert w.status_code == 200 and w.content[:2] == b"PK"
        lc = client.get(f"/api/subject-screens/{sid}/link-chart.xlsx")
        assert lc.status_code == 200 and "Links" in load_workbook(io.BytesIO(lc.content)).sheetnames
        assert client.get("/api/subject-screens/nope").status_code == 404
        assert any(a["action"] == "subject_screen" for a in client.get("/api/audit").json())
    finally:
        mp.undo()


def test_word_reports(syn):
    from docx import Document

    from ledgerhawk.exports.word import PRIVILEGED as WORD_PRIVILEGED, build_case_docx, build_subjects_docx
    sam, ex, p, _ = syn
    res = subject_screen(parse_subjects(f"{p['ex_affiliate']}\n{p['excluded_major']}"), sam, ex).to_dict()
    res["meta"] = {"id": "x", "matter": "M-1", "client": "Client", "privileged": True, "created_by": "A", "created_at": "2026-10-04"}
    doc = Document(io.BytesIO(build_subjects_docx(res)))
    text = "\n".join(par.text for par in doc.paragraphs)
    assert "What we found" in text and "Next steps" in text and "Method and limits" in text
    assert doc.sections[0].header.paragraphs[0].text == WORD_PRIVILEGED
    assert doc.core_properties.author == "LedgerHawk" and doc.core_properties.last_modified_by == "LedgerHawk"
    assert "python-docx" not in (doc.core_properties.comments or "")

    v = res["subjects"][0]["entities"][0] | {"queue": "exclusion", "naics": "", "psc": ""}
    summary = {"meta": {"id": "run"}, "manifest": {"sam_extract_date": "2026-09-06"}}
    doc = Document(io.BytesIO(build_case_docx(v, {"tier": "3"}, None, [], summary, matter="M-1")))
    text = "\n".join(par.text for par in doc.paragraphs)
    assert "Why it flagged" in text and "Investigator notes" in text and "Matter: M-1" in text
    assert not doc.sections[0].header.paragraphs[0].text


def _later_exclusions(excl, out, add_uei: str, drop_uei: str):
    """A later exclusions extract: one new firm exclusion, one record dropped."""
    import csv
    rows = list(csv.DictReader(open(excl)))
    rows = [r for r in rows if r["Unique Entity ID"] != drop_uei]
    rows.append({**rows[0], "Name": "", "Classification": "Firm", "Unique Entity ID": add_uei, "Excluding Agency": "DOL",
                 "Exclusion Type": "Ineligible (Proceedings Completed)", "Active Date": "10/05/2026", "Termination Date": "Indefinite",
                 "Additional Comments": ""})
    with open(out, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


def test_compare_screens(syn, tmp_path):
    from ledgerhawk.pipeline.subjects import compare_screens
    sam, ex, p, (_, excl, _) = syn
    subs = parse_subjects("\n".join([p["succ_old"], p["excluded_major"], p["ex_affiliate"]]))
    old = subject_screen(subs, sam, ex).to_dict() | {"meta": {"id": "old", "created_at": "2026-10-04T00:00:00"}}
    later = load_exclusions(_later_exclusions(excl, tmp_path / "ex2.csv", p["succ_old"], p["excluded_major"]), date(2026, 10, 9))
    new = subject_screen(subs, sam, later).to_dict()
    ch = compare_screens(old, new)
    rows = {r["ref"]: r for r in ch["subjects"]}
    assert rows[1]["direction"] == "worse" and rows[1]["status_now"] == "excluded"
    assert any("DOL" in x for x in rows[1]["added"])
    assert rows[2]["direction"] == "better" and any("GSA" in x for x in rows[2]["removed"])
    assert 3 not in rows and ch["counts"] == {"changed": 2, "worse": 1, "better": 1, "unchanged": 1}
    # The related entity (succ_new) sees nothing new: the change is on the subject itself.
    assert ch["subjects"][0]["ref"] == 1  # worse first


def test_recheck_api(syn, tmp_path):
    from docx import Document
    _, _, p, (vendors, excl, sam) = syn
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    client = TestClient(appmod.app)
    try:
        ids = {}
        for kind, path, d in (("sam", sam, "2026-09-06"), ("exclusions", excl, "2026-10-02")):
            with open(path, "rb") as f:
                ids[kind] = client.post("/api/sources", files={"file": f}, data={"kind": kind, "as_of": d, "analyst": "T"}).json()["id"]
        r = client.post("/api/subject-screens", data={"subjects_text": f"{p['succ_old']}\n{p['excluded_major']}", "analyst": "T",
                                                      "matter": "M-1", "sam_source": ids["sam"], "exclusions_source": ids["exclusions"]})
        first = r.json()["id"]
        later = _later_exclusions(excl, tmp_path / "ex2.csv", p["succ_old"], p["excluded_major"])
        with open(later, "rb") as f:
            client.post("/api/sources", files={"file": ("ex2.csv", f)}, data={"kind": "exclusions", "as_of": "2026-10-09", "analyst": "T"})
        assert client.post(f"/api/subject-screens/{first}/recheck", data={"analyst": ""}).status_code == 400
        assert client.post("/api/subject-screens/nope/recheck", data={"analyst": "T"}).status_code == 404
        r = client.post(f"/api/subject-screens/{first}/recheck", data={"analyst": "T"})
        assert r.status_code == 200, r.text
        assert r.json()["state"] == "done" and r.json()["kind"] == "recheck"
        got = client.get(f"/api/subject-screens/{r.json()['result']['id']}").json()
        assert got["meta"]["parent_id"] == first and got["meta"]["matter"] == "M-1"
        assert got["sources"]["exclusions_extract_date"] == "2026-10-09"  # newest extract picked up
        assert got["changes"]["counts"]["worse"] == 1
        wb = load_workbook(io.BytesIO(client.get(f"/api/subject-screens/{got['meta']['id']}/subject-screen.xlsx").content))
        assert "What Changed" in wb.sheetnames
        doc = Document(io.BytesIO(client.get(f"/api/subject-screens/{got['meta']['id']}/subject-screen.docx").content))
        assert any(par.text == "What changed since the last check" for par in doc.paragraphs)
        assert client.get("/api/subject-screens").json()[0]["change_counts"]["changed"] == 2

        # FY dollar columns on an uploaded list survive a re-check, so the data checks still have the file's figures
        lst = tmp_path / "voi.csv"
        lst.write_text(f"UEI,Name,FY24 Obligated ($),FY25 Obligated ($)\n{p['succ_old']},Old Co,\"$6,450,000\",0\n")
        with open(lst, "rb") as f:
            sid = client.post("/api/subject-screens", files={"subjects_file": ("voi.csv", f)},
                              data={"analyst": "T", "sam_source": ids["sam"], "exclusions_source": ids["exclusions"]}).json()["id"]
        again = client.post(f"/api/subject-screens/{sid}/recheck", data={"analyst": "T"}).json()["result"]["id"]
        ent = client.get(f"/api/subject-screens/{again}").json()["subjects"][0]["entities"][0]
        assert ent["fy24"] == 6450000.0 and ent["dollars_from"] == "list"
    finally:
        mp.undo()


def test_people(syn):
    from ledgerhawk.pipeline.subjects import parse_people, people_screen
    sam, ex, _, _ = syn
    people = parse_people("Reese Fosterling, NY\nDrew Excludedson, TX\nDrew Excludedson\nPat Synthetic, CA\nAvery Agentworth\nNobody Here")
    assert people[0] == {"input": "Reese Fosterling, NY", "first": "REESE", "last": "FOSTERLING", "state": "NY", "ref": 1}
    with pytest.raises(ValueError):
        parse_people("Cher")
    r = {x["ref"]: x for x in people_screen(people, sam, ex)}
    assert r[1]["status"] == "tied" and sum(x["excluded"] for x in r[1]["registrations"]) == 3
    assert r[2]["status"] == "excluded"
    assert r[3]["status"] == "name_only"  # no state given: a name alone never confirms
    assert r[4]["status"] == "name_only" and "different state (TX)" in r[4]["exclusions"][0]["support"]
    assert r[5]["status"] == "listed" and r[5]["common"]  # registered-agent contact on 9 registrations
    assert r[6]["status"] == "clear"


def test_people_only_screen_api(syn, tmp_path):
    from docx import Document
    _, _, p, (vendors, excl, sam) = syn
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    client = TestClient(appmod.app)
    try:
        ids = {}
        for kind, path, d in (("sam", sam, "2026-09-06"), ("exclusions", excl, "2026-10-02")):
            with open(path, "rb") as f:
                ids[kind] = client.post("/api/sources", files={"file": f}, data={"kind": kind, "as_of": d, "analyst": "T"}).json()["id"]
        base = {"analyst": "T", "sam_source": ids["sam"], "exclusions_source": ids["exclusions"]}
        assert client.post("/api/subject-screens", data={**base, "people_text": "Cher"}).status_code == 400
        r = client.post("/api/subject-screens", data={**base, "people_text": "Reese Fosterling, NY"})
        assert r.status_code == 200, r.text
        sid = r.json()["id"]
        got = client.get(f"/api/subject-screens/{sid}").json()
        assert got["subjects"] == [] and got["people"][0]["status"] == "tied" and got["counts"]["people"] == 1
        wb = load_workbook(io.BytesIO(client.get(f"/api/subject-screens/{sid}/subject-screen.xlsx").content))
        assert "People" in wb.sheetnames
        doc = Document(io.BytesIO(client.get(f"/api/subject-screens/{sid}/subject-screen.docx").content))
        assert any(par.text.startswith("Person 1.") for par in doc.paragraphs)
        r = client.post(f"/api/subject-screens/{sid}/recheck", data={"analyst": "T"})
        assert r.status_code == 200, r.text
        again = client.get(f"/api/subject-screens/{r.json()['result']['id']}").json()
        assert again["people"][0]["status"] == "tied" and again["changes"]["people"] == []
    finally:
        mp.undo()


def test_notes_evidence_and_signoff(syn, tmp_path):
    import hashlib
    from docx import Document
    _, _, p, (vendors, excl, sam) = syn
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    client = TestClient(appmod.app)
    try:
        ids = {}
        for kind, path, d in (("sam", sam, "2026-09-06"), ("exclusions", excl, "2026-10-02")):
            with open(path, "rb") as f:
                ids[kind] = client.post("/api/sources", files={"file": f}, data={"kind": kind, "as_of": d, "analyst": "T"}).json()["id"]
        sid = client.post("/api/subject-screens", data={"subjects_text": p["ex_affiliate"], "people_text": "Reese Fosterling, NY",
                                                        "analyst": "Ana", "matter": "M-9", "sam_source": ids["sam"],
                                                        "exclusions_source": ids["exclusions"]}).json()["id"]
        base = f"/api/subject-screens/{sid}"
        assert client.get(base).json()["review"]["state"] == "draft"
        assert client.post(f"{base}/notes", data={"analyst": "", "text": "x"}).status_code == 400
        assert client.post(f"{base}/notes", data={"analyst": "Ana", "target": "s:9", "text": "x"}).status_code == 400
        assert client.post(f"{base}/notes", data={"analyst": "Ana", "target": "s:1"}).status_code == 400
        pdf = b"%PDF-1.4 sam record"
        r = client.post(f"{base}/notes", data={"analyst": "Ana", "target": "s:1", "text": "Called the registered POC; number disconnected.",
                                               "source": "SAM.gov entity record, viewed 2026-10-04"},
                        files={"file": ("SAM record (1).pdf", pdf, "application/pdf")})
        assert r.status_code == 200, r.text
        note = r.json()
        assert note["file"] == "SAM_record_1_.pdf" and note["file_sha256"] == hashlib.sha256(pdf).hexdigest()
        assert client.get(f"{base}/evidence/{note['id']}").content == pdf
        other = client.post(f"{base}/notes", data={"analyst": "Ana", "target": "p:1", "text": "Same person per state filing."}).json()
        gone = client.post(f"{base}/notes", data={"analyst": "Ana", "target": "screen", "text": "typo"}).json()
        assert client.post(f"{base}/notes/{gone['id']}/delete", data={"analyst": "Rev"}).status_code == 400  # not the author
        assert client.post(f"{base}/notes/{gone['id']}/delete", data={"analyst": " ana "}).status_code == 200
        assert [n["id"] for n in client.get(base).json()["review"]["notes"]] == [note["id"], other["id"]]

        # two-person sign-off
        assert client.post(f"{base}/review", data={"analyst": "Rev", "action": "approve"}).status_code == 400  # not submitted
        assert client.post(f"{base}/review", data={"analyst": "Ana", "action": "submit"}).status_code == 200
        assert client.post(f"{base}/review", data={"analyst": "ana", "action": "approve"}).status_code == 400  # same person
        assert client.post(f"{base}/review", data={"analyst": "Rev", "action": "return"}).status_code == 400  # needs a comment
        assert client.post(f"{base}/review", data={"analyst": "Rev", "action": "return", "comment": "Cite the state filing."}).json()["state"] == "returned"
        client.post(f"{base}/review", data={"analyst": "Ana", "action": "submit"})
        r = client.post(f"{base}/review", data={"analyst": "Rev", "action": "approve"})
        assert r.json()["state"] == "approved" and [h["action"] for h in r.json()["history"]] == ["submit", "return", "submit", "approve"]
        assert client.post(f"{base}/notes", data={"analyst": "Ana", "target": "screen", "text": "late"}).status_code == 400  # locked
        assert client.get("/api/subject-screens").json()[0]["review_state"] == "approved"

        doc = Document(io.BytesIO(client.get(f"{base}/subject-screen.docx").content))
        text = "\n".join(par.text for par in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        assert "Called the registered POC" in text and "Same person per state filing." in text and "typo" not in text
        assert "Approved by" in text and "Rev" in text and note["file_sha256"] in text
        wb = load_workbook(io.BytesIO(client.get(f"{base}/subject-screen.xlsx").content))
        assert wb["Analyst Notes"].max_row == 7 and wb["Analyst Notes"]["C6"].value.startswith("Called")
        assert any(a["action"] == "screen_approve" for a in client.get("/api/audit").json())

        # a re-check carries notes and evidence forward and starts sign-off again
        new = client.post(f"{base}/recheck", data={"analyst": "Ana"}).json()["result"]["id"]
        rv = client.get(f"/api/subject-screens/{new}").json()["review"]
        assert rv["state"] == "draft" and len(rv["notes"]) == 2 and rv["notes"][0]["carried_from"]["id"] == sid
        assert client.get(f"/api/subject-screens/{new}/evidence/{note['id']}").content == pdf
        assert client.post(f"{base}/review", data={"analyst": "Rev", "action": "reopen"}).status_code == 400  # needs a comment
        assert client.post(f"{base}/review", data={"analyst": "Rev", "action": "reopen", "comment": "New filing"}).json()["state"] == "draft"
    finally:
        mp.undo()


def test_link_chart(syn, tmp_path):
    from ledgerhawk.exports.linkchart import build_chart, build_linkchart
    from ledgerhawk.pipeline.subjects import parse_people, people_screen
    sam, ex, p, _ = syn
    res = subject_screen(parse_subjects(f"{p['ex_affiliate']}\n{p['excluded_major']}"), sam, ex).to_dict()
    res["people"] = people_screen(parse_people("Reese Fosterling, NY"), sam, ex)
    res["meta"] = {"id": "x", "matter": "M", "privileged": True, "data_class": "synthetic"}
    ch = build_chart(res)
    ents, links = ch.entities, ch.links
    sub = next(s for s in res["subjects"] if s["input_uei"] == p["ex_affiliate"] or s["input_name"] == p["ex_affiliate"])
    org = f"ORG:{sub['entities'][0]['uei']}"
    assert ents[org]["subjects"] == {sub["ref"]}
    # the shared contact and suite become entities that related excluded firms also link to
    contact = next(x["to"] for x in links if x["from"] == org and x["type"] == "Has SAM contact")
    suite = next(x["to"] for x in links if x["from"] == org and x["type"] == "Registered at")
    rel = [r for r in sub["related"] if r["excluded"]]
    assert rel and all({"from": f"ORG:{r['uei']}", "to": contact} in [{"from": x["from"], "to": x["to"]} for x in links] for r in rel)
    assert any(x["from"] == f"ORG:{rel[0]['uei']}" and x["to"] == suite for x in links)
    assert any(x["type"] == "Excluded under this UEI" for x in links)
    assert contact == "PER:REESE|FOSTERLING|NY" and "Screened person 1" in ents[contact]["description"]  # screened person merges with the SAM contact
    assert all(x["from"] in ents and x["to"] in ents for x in links)
    wb = load_workbook(io.BytesIO(build_linkchart(res)))
    assert wb.sheetnames == ["Read Me", "Entities", "Links"]
    assert wb["Read Me"]["A1"].value == PRIVILEGED
    assert [c.value for c in wb["Links"][1]][:4] == ["From ID", "From Type", "From Label", "Link Type"]
    assert wb["Links"].max_row == len(links) + 1 and wb["Entities"].max_row == len(ents) + 1


def _fake_usaspending(calls):
    """Stands in for api.usaspending.gov: two contracts for every UEI, one IDV, and a stray award to another UEI."""
    import urllib.error

    def post(url, body):
        calls.append(body)
        uei = body["filters"]["recipient_search_text"][0]
        if uei == "FAILFAILFAIL":
            raise urllib.error.URLError("down")
        if url.endswith("/spending_over_time/"):  # FY23 small, FY25 a jump
            return {"results": [{"time_period": {"fiscal_year": 2023}, "aggregated_amount": 100000, "Contract_Obligations": 80000,
                                 "Idv_Obligations": 20000},
                                {"time_period": {"fiscal_year": 2025}, "aggregated_amount": 2.5e6, "Contract_Obligations": 2.5e6,
                                 "Idv_Obligations": None}]}
        if url.endswith("/spending_by_transaction/"):
            if "IDV_A" in body["filters"]["award_type_codes"]:
                # FPDS codes this PO mod N (legal contract cancellation); a PO mod with no code reads as an option
                rows = [{"Award ID": "47QTCA24D003M", "Mod": "PO0010", "Action Date": "2026-06-12", "Action Type": "N",
                         "Transaction Amount": 0, "Recipient UEI": uei, "Awarding Agency": "General Services Administration"},
                        {"Award ID": "47QTCA24D003M", "Mod": "PO0011", "Action Date": "2026-07-01", "Action Type": None,
                         "Transaction Amount": 0, "Recipient UEI": uei, "Awarding Agency": "General Services Administration"}]
            else:
                rows = [{"Award ID": "47QTCA25F0001", "Mod": "0", "Action Date": "2025-07-01", "Action Type": None,
                         "Transaction Amount": 250000.5, "Recipient UEI": uei, "Awarding Agency": "General Services Administration"},
                        {"Award ID": "47QTCA22F0002", "Mod": "P00004", "Action Date": "2025-08-01", "Action Type": "C",
                         "Transaction Amount": -1200, "Recipient UEI": uei},  # deobligation: not flagged
                        {"Award ID": "47QTCA22F0002", "Mod": "P00005", "Action Date": "2025-09-01", "Action Type": "K",
                         "Transaction Amount": 0, "Recipient UEI": uei},  # closeout: not flagged
                        {"Award ID": "OLD", "Mod": "P00001", "Action Date": "2024-01-01", "Action Type": "G",
                         "Transaction Amount": 5000, "Recipient UEI": uei}]  # before the exclusion: dropped
            return {"results": rows, "page_metadata": {"page": 1, "hasNext": False}}
        if "IDV_A" in body["filters"]["award_type_codes"]:
            return {"results": [{"internal_id": 3, "generated_internal_id": f"CONT_IDV_{uei}", "Award ID": "GS-00F-001",
                                 "Recipient UEI": uei, "Award Amount": 0, "Awarding Agency": "General Services Administration",
                                 "Start Date": "2019-03-01", "Last Date to Order": "2029-02-28", "Contract Award Type": "FSS"}],
                    "page_metadata": {"page": 1, "hasNext": False}}
        return {"results": [
            {"internal_id": 1, "generated_internal_id": f"CONT_AWD_A_{uei}", "Award ID": "47QTCA25F0001", "Recipient UEI": uei,
             "Recipient Name": "X", "Award Amount": 250000.5, "Awarding Agency": "General Services Administration",
             "Awarding Sub Agency": "Federal Acquisition Service", "Start Date": "2025-07-01", "End Date": "2026-06-30",
             "Contract Award Type": "DELIVERY ORDER", "NAICS": {"code": "541611", "description": "Admin consulting"}, "PSC": "R408"},
            {"internal_id": 2, "generated_internal_id": f"CONT_AWD_B_{uei}", "Award ID": "47QTCA22F0002", "Recipient UEI": uei,
             "Award Amount": 90000, "Awarding Agency": "Department of Veterans Affairs", "Start Date": "2022-01-15"},
            {"internal_id": 9, "Award ID": "OTHER", "Recipient UEI": "ZZZZZZZZZZZZ", "Award Amount": 5e6, "Start Date": "2024-01-01"},
        ], "page_metadata": {"page": 1, "hasNext": True}}
    return post


def test_awards_lookup(syn):
    from datetime import date as d
    from ledgerhawk.pipeline.awards import awards_for_uei, screen_awards
    sam, ex, p, _ = syn
    calls = []
    one = awards_for_uei("ABCDEFGHJKLM", _fake_usaspending(calls), today=d(2026, 10, 4))
    assert [a["award_id"] for a in one["awards"]] == ["47QTCA25F0001", "47QTCA22F0002", "GS-00F-001"]  # stray UEI dropped
    assert one["truncated"] and one["awards"][0]["naics"] == "541611" and one["awards"][0]["url"].endswith("CONT_AWD_A_ABCDEFGHJKLM")
    assert one["awards"][2]["end"] == "2029-02-28"  # IDVs report a last date to order
    assert {tuple(c["filters"]["award_type_codes"][:1]) for c in calls} == {("A",), ("IDV_A",)}  # groups asked separately
    assert calls[0]["filters"]["time_period"][0]["end_date"] == "2026-10-04"
    assert awards_for_uei("FAILFAILFAIL", _fake_usaspending([]))["error"].startswith("USAspending did not answer")

    res = subject_screen(parse_subjects(f"{p['excluded_major']}\n{p['ex_affiliate']}"), sam, ex).to_dict()
    aw = screen_awards(res, _fake_usaspending([]), today=d(2026, 10, 4))
    by = {e["uei"]: e for e in aw["entities"]}
    major = by[p["excluded_major"]]
    assert major["role"] == "subject" and major["excluded_since"] == "2025-03-01"
    assert major["after_exclusion"] == 1 and major["total"] == 340000.5  # the July 2025 order came after the exclusion
    assert any(e["role"] == "related, excluded" for e in aw["entities"])  # excluded related firms are looked up too
    affiliate = next(e for e in aw["entities"] if e["role"] == "subject" and e["uei"] != p["excluded_major"])
    assert affiliate["after_exclusion"] == 0 and not affiliate["excluded_since"]
    assert not affiliate["actions"]  # no exclusion, no transaction lookup

    # every action after the exclusion, by kind: a new order and an uncoded GSA Schedule PO mod are flagged; the FPDS
    # code wins over the PO numbering, so PO0010 (coded N) is a cancellation, said with how long it took
    kinds = [(a["award_id"], a["mod"], a["kind"], a["flagged"]) for a in major["actions"]]
    assert kinds == [("47QTCA25F0001", "0", "new", True), ("47QTCA22F0002", "P00004", "admin", False),
                     ("47QTCA22F0002", "P00005", "wind_down", False), ("47QTCA24D003M", "PO0010", "cancel", False),
                     ("47QTCA24D003M", "PO0011", "option", True)]
    assert major["actions_flagged"] == 2 and major["schedule_actions"] == 2 and major["actions_dollars"] == 250000.5
    assert "GSA Schedule 47QTCA24D003M was modified after the exclusion (PO0011, 2026-07-01)" in major["actions_summary"]
    assert ("GSA Schedule 47QTCA24D003M was cancelled or terminated PO0010 on 2026-06-12, 468 days after the exclusion "
            "(FPDS action type N)") in major["actions_summary"]
    assert major["by_fy"] == {"2023": 100000.0, "2025": 2500000.0} and major["lifetime"] == 2600000.0
    assert major["growth"].startswith("FY25 obligations of $2.5M are 25x its best earlier year (FY23, $100,000")


def test_growth_and_money_shift_notes():
    from datetime import date as d
    from ledgerhawk.pipeline.awards import growth_note, shift_note
    today = d(2026, 10, 6)  # FY27 has just started
    assert growth_note({2020: 300000, 2026: 400000}, today) == ""  # under $1M is not a story
    assert growth_note({2020: 300000, 2026: 1.2e6}, today) == ""  # 4x is ordinary growth
    assert "8x its best earlier year" in growth_note({2019: 150000, 2020: 200000, 2026: 1.6e6}, today)
    assert growth_note({2026: 3e6}, today) == "No federal contract dollars before FY26; $3.0M in FY26"
    # a second big year is measured against the record before the run-up, not against the first big year (ELB)
    assert growth_note({2022: 47000, 2025: 28.5e6, 2026: 32.1e6}, today).startswith(
        "FY26 obligations of $32.1M are 683x its best earlier year (FY22, $47,000")
    assert growth_note({2025: 2e6, 2026: 2.5e6}, today) == "No federal contract dollars before FY25; $2.5M in FY26"
    shift = shift_note([{"uei": "OLDUEI", "by_fy": {2024: 1e6, 2025: 5e6, 2026: 400000}},
                        {"uei": "NEWUEI", "by_fy": {2026: 4.2e6}}], today)
    assert shift.startswith("Money moved between registrations: as OLDUEI fell from $5.0M in FY25 to $400,000 in FY26, NEWUEI went")
    assert shift_note([{"uei": "A", "by_fy": {2025: 5e6, 2026: 5e6}}, {"uei": "B", "by_fy": {2026: 4e6}}], today) == ""  # A did not fall
    # early in FY27 the FY24 -> FY25 move still counts (K2: Manassas fell as Wrightstown rose from nothing)
    k2 = shift_note([{"uei": "KCYXLTWPA849", "by_fy": {2023: 17.2e6, 2024: 15.0e6, 2025: 5.47e6, 2026: 3.06e6}},
                     {"uei": "LLMJLD2G7RD5", "by_fy": {2023: 3.27e6, 2025: 27.4e6, 2026: 3.31e6}}], today)
    assert k2 == ("Money moved between registrations: as KCYXLTWPA849 fell from $15.0M in FY24 to $5.5M in FY25, "
                  "LLMJLD2G7RD5 went from $0 to $27.4M")


def test_awards_api_and_exports(syn, tmp_path):
    from docx import Document
    _, _, p, (vendors, excl, sam) = syn
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    client = TestClient(appmod.app)
    try:
        ids = {}
        for kind, path, d in (("sam", sam, "2026-09-06"), ("exclusions", excl, "2026-10-02")):
            with open(path, "rb") as f:
                ids[kind] = client.post("/api/sources", files={"file": f}, data={"kind": kind, "as_of": d, "analyst": "T"}).json()["id"]
        sid = client.post("/api/subject-screens", data={"subjects_text": p["excluded_major"], "analyst": "T",
                                                        "sam_source": ids["sam"], "exclusions_source": ids["exclusions"]}).json()["id"]
        base = f"/api/subject-screens/{sid}"
        assert client.get(base).json()["awards"] is None
        appmod.store.awards_post = lambda url, body: (_ for _ in ()).throw(OSError("blocked"))
        r = client.post(f"{base}/awards", data={"analyst": "T"})
        assert r.status_code == 200 and r.json()["state"] == "error" and "USAspending did not answer" in r.json()["error"]
        assert client.get(base).json()["awards"] is None  # a failed lookup saves nothing
        assert client.get(f"{base}/jobs").json()["awards"]["state"] == "error"
        appmod.store.awards_post = _fake_usaspending([])
        assert client.post(f"{base}/awards", data={"analyst": ""}).status_code == 400
        r = client.post(f"{base}/awards", data={"analyst": "Ana"})
        assert r.status_code == 200, r.text
        job = r.json()
        assert job["state"] == "done" and job["by"] == "Ana" and job["result"]["ueis"] >= 1 and job["total"] == 4 * job["result"]["ueis"]
        assert job["done"] == job["total"]  # progress reached the end
        got = client.get(base).json()["awards"]
        assert got["fetched_by"] == "Ana" and got["entities"][0]["after_exclusion"] == 1
        wb = load_workbook(io.BytesIO(client.get(f"{base}/subject-screen.xlsx").content))
        ws = wb["Awards"]
        assert ws["F6"].value == "47QTCA25F0001" and ws["M6"].value == "Yes" and ws["L6"].value == 250000.5
        doc = Document(io.BytesIO(client.get(f"{base}/subject-screen.docx").content))
        text = "\n".join(par.text for par in doc.paragraphs)
        assert "Federal awards (USAspending)" in text and "started on or after the exclusion of 2025-03-01" in text
        assert "GSA Schedule 47QTCA24D003M was modified after the exclusion (PO0011, 2026-07-01)" in text
        ws = wb["After Exclusion"]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=6, min_col=2, max_col=12)]
        assert ["47QTCA24D003M", "PO0010", "cancelled or terminated", None, "Yes"] == rows[-2][5:10]
        assert ["47QTCA24D003M", "PO0011", "option exercised", "Yes", "Yes"] == rows[-1][5:10]
        assert any(t.cell(0, 0).text == "Date" for t in doc.tables)  # the after-exclusion table in the report
        assert "25x its best earlier year" in text
        ws = wb["By Fiscal Year"]
        assert [c.value for c in ws[5]][1:8] == ["Subject #", "UEI", "Entity", "Role", "FY23 ($)", "FY25 ($)", "Lifetime ($)"]
        assert ws["F6"].value == 100000 and ws["H6"].value == 2600000 and ws["I6"].value.startswith("FY25 obligations")
        assert any(a["action"] == "screen_awards" for a in client.get("/api/audit").json())
    finally:
        mp.undo()


def test_subject_file_with_title_rows_above_the_header(tmp_path):
    # hand-built lead lists put a title and notes above the column headers
    p = tmp_path / "voi.xlsx"
    rows = [["LedgerHawk GSA Vendors of Interest", None, None], ["Built 2026-10-04", None, None], [None, None, None],
            ["#", "UEI", "Name"], [1, "HFGCD12199B8", "ELB SERVICES LLC"], [2, "", "K2 CONTRACTING GROUP LLC"]]
    pd.DataFrame(rows).to_excel(p, header=False, index=False)
    got = parse_subjects(path=p)
    assert [(s["uei"], s["name"]) for s in got] == [("HFGCD12199B8", "ELB SERVICES LLC"), ("", "K2 CONTRACTING GROUP LLC")]


def test_subject_list_dollars_and_data_checks(tmp_path):
    from ledgerhawk.pipeline.awards import file_mismatch, odd_awards
    p = tmp_path / "voi.csv"
    p.write_text('UEI,Name,FY24 Obligated ($),FY25 Obligated ($)\nHFGCD12199B8,MUCOMMUNE LLC,"$6,450,000",(1200)\n')
    got = parse_subjects(path=p)
    assert got[0]["fy24"] == 6450000.0 and got[0]["fy25"] == -1200.0  # the list's own dollar columns are kept
    assert file_mismatch({2024: 6.45e6, 2025: 0}, {2024: 306000}) == [
        "Data check: the GSA file shows $6.5M for FY24; USAspending shows $306,000. Confirm the record before it counts toward totals"]
    assert file_mismatch({2024: 1.2e6}, {2024: 900000}) == []  # close enough
    odd = odd_awards([{"award_id": "19AQMM24P0001", "amount": 7.79e6, "start": "2024-03-01", "end": "2024-03-31"},
                      {"award_id": "LONG", "amount": 9e6, "start": "2024-01-01", "end": "2025-01-01"}])
    assert odd == ["Data check: 19AQMM24P0001 is $7.8M for 31 days of work (2024-03-01 to 2024-03-31). "
                   "The amount may be in local currency or mis-keyed"]

    # a short span that is really a termination (ATI's IRS contract, terminated for convenience after 42 days)
    def fpds(url, body):
        rows = {"2023H225C00025": [{"Award ID": "2023H225C00025", "Mod": "0", "Action Date": "2025-01-31", "Action Type": None},
                                   {"Award ID": "2023H225C00025", "Mod": "P00001", "Action Date": "2025-03-14", "Action Type": "F"}],
                "19AQMM24P0001": [{"Award ID": "19AQMM24P0001", "Mod": "P00001", "Action Date": "2024-06-12", "Action Type": "C"}]}
        return {"results": rows[body["filters"]["award_ids"][0]]}
    odd = odd_awards([{"group": "contract", "award_id": "2023H225C00025", "amount": 19787636.48, "start": "2025-01-31",
                       "end": "2025-03-14"},
                      {"group": "contract", "award_id": "19AQMM24P0001", "amount": 7.79e6, "start": "2024-03-01",
                       "end": "2024-03-31"}], fpds)
    assert odd[0] == ("Data check: 2023H225C00025 shows $19.8M obligated, but FPDS records a termination or cancellation "
                      "(P00001, action type F, 2025-03-14) 42 days after award, and the amount was not brought down. "
                      "Confirm whether the money was deobligated")
    assert odd[1].endswith("The amount may be in local currency or mis-keyed")  # no termination: the H&S Water case


def _screen_app(syn, tmp_path, subjects_text):
    """A fresh app with the synthetic extracts loaded and one screen of `subjects_text`."""
    _, _, p, (vendors, excl, sam) = syn
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    client = TestClient(appmod.app)
    ids = {}
    for kind, path, d in (("sam", sam, "2026-09-06"), ("exclusions", excl, "2026-10-02")):
        with open(path, "rb") as f:
            ids[kind] = client.post("/api/sources", files={"file": f}, data={"kind": kind, "as_of": d, "analyst": "T"}).json()["id"]
    sid = client.post("/api/subject-screens", data={"subjects_text": subjects_text, "analyst": "T", "matter": "M-1",
                                                    "sam_source": ids["sam"], "exclusions_source": ids["exclusions"]}).json()["id"]
    return mp, appmod, client, sid


def test_screen_jobs_run_in_the_background_with_progress(syn, tmp_path):
    import threading
    import time
    _, _, p, _ = syn
    mp, appmod, client, sid = _screen_app(syn, tmp_path, f"{p['excluded_major']}\n{p['ex_affiliate']}")
    try:
        appmod.store.jobs_inline = False
        gate = threading.Event()
        fake = _fake_usaspending([])

        def slow(url, body):  # hold the lookup until the test has seen it running
            gate.wait(10)
            return fake(url, body)
        appmod.store.awards_post = slow
        base = f"/api/subject-screens/{sid}"
        first = client.post(f"{base}/awards", data={"analyst": "Ana"}).json()
        assert first["state"] == "running" and first["label"] == "Looking up awards on USAspending" and first["by"] == "Ana"
        again = client.post(f"{base}/awards", data={"analyst": "Bo"}).json()  # a second click joins the running job
        assert again["started_at"] == first["started_at"] and again["by"] == "Ana"
        assert client.get(base).json()["jobs"]["awards"]["state"] == "running"  # the page sees it on reload
        gate.set()
        seen = []
        for _ in range(200):
            j = client.get(f"{base}/jobs").json()["awards"]
            seen.append(j["done"])
            if j["state"] != "running":
                break
            time.sleep(0.05)
        assert j["state"] == "done" and j["done"] == j["total"] > 0 and j["step"] == "Done" and j["finished_at"]
        assert j["result"]["ueis"] * 4 == j["total"] and seen == sorted(seen)  # progress only moves forward
        assert client.get(base).json()["awards"]["fetched_by"] == "Ana"
        assert client.get("/api/subject-screens/nope/jobs").status_code == 404
        assert client.post("/api/subject-screens/nope/awards", data={"analyst": "A"}).status_code == 404
    finally:
        mp.undo()


def test_a_job_cut_off_by_a_restart_says_so(syn, tmp_path):
    import json
    _, _, p, _ = syn
    mp, appmod, client, sid = _screen_app(syn, tmp_path, p["excluded_major"])
    try:
        f = tmp_path / "data" / "subjects" / sid / "jobs.json"
        f.write_text(json.dumps({"context": {"kind": "context", "state": "running", "done": 3, "total": 40}}))
        j = client.get(f"/api/subject-screens/{sid}/jobs").json()["context"]
        assert j["state"] == "error" and "server restarted" in j["error"]
    finally:
        mp.undo()


def test_rename_a_screen(syn, tmp_path):
    from docx import Document
    _, _, p, _ = syn
    mp, appmod, client, sid = _screen_app(syn, tmp_path, p["excluded_major"])
    base = f"/api/subject-screens/{sid}"
    try:
        assert client.post(f"{base}/rename", data={"analyst": "", "matter": "X"}).status_code == 400
        assert client.post("/api/subject-screens/nope/rename", data={"analyst": "A", "matter": "X"}).status_code == 404
        assert client.post(f"{base}/rename", data={"analyst": "A", "matter": "x" * 201}).status_code == 400
        meta = client.post(f"{base}/rename", data={"analyst": "Ana", "matter": "GSA Vendors of Interest", "client": "GSA FAS"}).json()
        assert meta["matter"] == "GSA Vendors of Interest" and meta["client"] == "GSA FAS" and meta["original_matter"] == "M-1"
        assert meta["renamed_by"] == "Ana"
        assert client.get("/api/subject-screens").json()[0]["matter"] == "GSA Vendors of Interest"
        audit = client.get("/api/audit").json()[0]
        assert audit["action"] == "subject_screen_renamed" and '"M-1" to "GSA Vendors of Interest"' in audit["detail"]
        doc = Document(io.BytesIO(client.get(f"{base}/subject-screen.docx").content))
        assert any("GSA Vendors of Interest" in par.text for par in doc.paragraphs)  # exports use the new name
        new = client.post(f"{base}/recheck", data={"analyst": "Ana"}).json()["result"]["id"]
        assert client.get(f"/api/subject-screens/{new}").json()["meta"]["matter"] == "GSA Vendors of Interest"  # carried forward
        # approval locks the name, as it locks the notes
        client.post(f"{base}/notes", data={"analyst": "Ana", "target": "screen", "text": "Checked."})
        client.post(f"{base}/review", data={"analyst": "Ana", "action": "submit"})
        client.post(f"{base}/review", data={"analyst": "Rev", "action": "approve"})
        assert client.get(base).json()["review"]["state"] == "approved"
        r = client.post(f"{base}/rename", data={"analyst": "Ana", "matter": "Other"})
        assert r.status_code == 400 and "locked" in r.json()["detail"]
    finally:
        mp.undo()


def test_outside_context_covers_every_subject(syn, tmp_path, monkeypatch):
    # the old 40-name cap is gone: a screen's lookup covers every subject entity and person
    from ledgerhawk.pipeline import context as context_mod
    from ledgerhawk.api.store import Store
    seen = {}

    def fake_many(targets, fetch=None, cache_dir=None, progress=None):
        seen["n"] = len(targets)
        for i, t in enumerate(targets, 1):
            progress and progress(i, len(targets), t["name"])
        return [{"name": t["name"], "uei": t.get("uei", ""), "person": bool(t.get("person")), "count": 0, "errors": 0,
                 "sources": {"x": {"items": [], "error": ""}}, "tally": {"strong": 0, "possible": 0, "weak": 0},
                 "adverse": 0, "fetched_at": "2026-10-06T00:00:00+00:00"} for t in targets]
    monkeypatch.setattr(context_mod, "lookup_many", fake_many)
    st = Store(tmp_path / "data")
    screen = {"subjects": [{"ref": n, "input_name": f"Firm {n}", "related": [],
                            "entities": [{"uei": f"UEI{n:09d}", "name": f"Firm {n} LLC", "sam": None}]} for n in range(1, 61)],
              "people": [{"ref": 1, "first": "Reese", "last": "Fosterling", "state": "NY", "registrations": []}]}
    monkeypatch.setattr(st, "subject_screen", lambda sid: screen)
    monkeypatch.setattr(st, "_save_context", lambda f: None)
    monkeypatch.setattr(st, "_with_verdicts", lambda f: f)
    (tmp_path / "data" / "subjects" / "S1").mkdir(parents=True)
    ticks = []
    st.screen_context("S1", "A", progress=lambda d, t, s: ticks.append((d, t)))
    assert seen["n"] == 61 and ticks[-1] == (61, 61)


def test_lookup_many_reports_progress_from_the_start(monkeypatch):
    from ledgerhawk.pipeline import context as context_mod
    monkeypatch.setattr(context_mod, "lookup", lambda name, **kw: {"name": name})
    calls = []
    out = context_mod.lookup_many([{"name": n} for n in ("A", "B", "C")], progress=lambda d, t, s: calls.append((d, t, s)))
    assert [o["name"] for o in out] == ["A", "B", "C"]
    assert calls[0][:2] == (0, 3) and calls[0][2].startswith("Looking up 3 names")  # the total shows before any name ends
    assert sorted(c[0] for c in calls[1:]) == [1, 2, 3] and all(c[2].startswith("Finished: ") for c in calls[1:])


def test_vendor_record_across_runs_and_screens(syn, tmp_path):
    _, _, p, (vendors, excl, sam) = syn
    mp, appmod, client, sid = _screen_app(syn, tmp_path, f"{p['excluded_major']}\n{p['ex_affiliate']}")
    try:
        appmod.store.awards_post = _fake_usaspending([])
        assert client.post(f"/api/subject-screens/{sid}/awards", data={"analyst": "A"}).json()["state"] == "done"
        srcs = {s["kind"]: s["id"] for s in client.get("/api/sources").json()["sources"]}
        with open(vendors, "rb") as v:
            run_id = client.post("/api/runs", files={"vendors": v}, data={"synthetic": "true", "analyst": "A", "sam_source": srcs["sam"],
                                                                           "exclusions_source": srcs["exclusions"]}).json()["id"]
        uei = p["excluded_major"]
        where = client.get(f"/api/vendors/{uei.lower()}").json()  # any case
        assert where["uei"] == uei and [r["id"] for r in where["runs"]] == [run_id] and where["name"]
        scr = where["screens"][0]
        assert scr["id"] == sid and scr["role"] == "subject" and scr["matter"] == "M-1" and scr["findings"]
        assert any(e["uei"] == uei for e in scr["awards"]["by_uei"]) and scr["awards"]["growth"].startswith("FY25 obligations")
        assert (tmp_path / "data" / "runs" / run_id / "names.json").exists()  # the lookup index is kept beside the run
        rel = next(r for s in client.get(f"/api/subject-screens/{sid}").json()["subjects"] for r in s["related"])
        rel_where = client.get(f"/api/vendors/{rel['uei']}").json()
        assert any(x["role"] == "related" and x["via"] for x in rel_where["screens"])
        assert client.get("/api/vendors/ZZZZZZZZZZZZ").status_code == 404
        # the run's vendor page carries the screens' evidence
        v = client.get(f"/api/runs/{run_id}/vendors/{uei}").json()
        assert v["screens"][0]["id"] == sid and v["screens"][0]["status"]
        # search by name or UEI, across runs and screens
        name = where["name"]
        hits = client.get("/api/vendors", params={"q": name.split()[0].lower()}).json()["rows"]
        mine = next(h for h in hits if h["uei"] == uei)
        assert mine["run"]["id"] == run_id and mine["runs"] == 1 and mine["screens"] == 1
        assert client.get("/api/vendors", params={"q": uei}).json()["rows"][0]["uei"] == uei
        assert client.get("/api/vendors", params={"q": "x"}).json()["rows"] == []
        # the link map carries the screens' related firms and money by fiscal year
        g = client.get(f"/api/runs/{run_id}/vendors/{uei}/graph").json()
        center = next(n for n in g["nodes"] if n.get("center"))
        assert center["in_run"] and center["money"] and set(center["money"]) <= set(g["years"]) and "Subject screen" in center["source"]
        scr_edges = [e for e in g["edges"] if e["kind"] == "screen_related"]
        for r in scr["related"]:
            assert any(r["uei"] in (e["source"][2:], e["target"][2:]) for e in scr_edges)
        assert all(n["source"] for n in g["nodes"])
        rg = client.get(f"/api/runs/{run_id}/vendors/{rel['uei']}/graph")
        if rg.status_code == 200:  # a related firm in the run links back to the subject it was screened with
            assert any(e["kind"] == "screen_related" and e["target"] == f"v:{rel['uei']}" for e in rg.json()["edges"])
        # a rename shows in the record
        client.post(f"/api/subject-screens/{sid}/rename", data={"analyst": "A", "matter": "Renamed"})
        assert client.get(f"/api/vendors/{uei}").json()["screens"][0]["matter"] == "Renamed"
    finally:
        mp.undo()
