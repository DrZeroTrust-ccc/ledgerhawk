"""The findings briefing: facts graded by what a record proves, cases ranked by their strongest fact, and a deck that
carries a source and a way to verify every fact."""
import importlib
import io
import json
from datetime import date

from fastapi.testclient import TestClient
from pptx import Presentation

from ledgerhawk.exports.findings import PAID_AFTER, PATTERN, SINGLE, build_cases, categorize, facts_for
from ledgerhawk.pipeline.synthetic import make_synthetic

MAN = {"exclusions_extract_date": "2026-10-02", "sam_extract_date": "2026-10-05", "input_file": "gsa.xlsx",
       "input_sha256": "abc123def4567890"}
EXCLUDED = {"uei": "U1", "name": "ACME LLC", "tot": 3e6, "fy24": 1e6, "fy25": 2e6, "exclusion_flags": ["EXCLUDED"],
            "exclusion": [{"kind": "direct", "active_date": "2025-06-18", "agency": "GSA", "type": "Ineligible (Proceedings Completed)",
                           "ct_code": "Z1"}]}
ENTITY = {"excluded_since": "2025-06-18", "fetched_at": "2026-10-10",
          "actions": [{"award_id": "47QSWA18D0001", "mod": "P00012", "date": "2025-09-01", "amount": 412000.0, "flagged": True,
                       "label": "work added or changed", "schedule": True, "url": "https://www.usaspending.gov/award/CONT_IDV_1"},
                      {"award_id": "47QSWA18D0001", "mod": "P00013", "date": "2025-06-18", "amount": 5000.0, "flagged": True,
                       "label": "funding added", "schedule": True, "url": ""}],
          "awards": [], "subawards": []}


def test_a_paid_action_after_exclusion_is_a_confirmed_fact_with_its_record_and_link():
    facts = facts_for(EXCLUDED, ENTITY, {"fetched_at": "2026-10-10", "excluded_since": "2025-06-18"}, [], MAN)
    confirmed = [f for f in facts if f["level"] == "confirmed"]
    assert confirmed[0]["text"].startswith("On the SAM exclusions list") and confirmed[0]["record"] == "CT code Z1"
    paid = [f for f in confirmed if f["kind"] == "paid_after"]
    assert len(paid) == 1  # the action on the exclusion date itself is not money after it
    assert paid[0]["record"] == "47QSWA18D0001 mod P00012" and paid[0]["url"].endswith("CONT_IDV_1")
    assert "(GSA Schedule)" in paid[0]["text"] and paid[0]["as_of"] == "2026-10-10"
    # with money after exclusion and no FAR record named, the contract file question is added
    assert any(f["level"] == "investigate" and "contracting officer" in f["text"] for f in facts)
    assert categorize(facts, "after") == PAID_AFTER
    assert [f["level"] for f in facts] == sorted([f["level"] for f in facts],
                                                 key=["confirmed", "anomaly", "indicator", "investigate"].index)


def test_indicators_name_an_innocent_explanation_and_name_only_matches_stay_weak():
    v = {"uei": "U2", "name": "BETA INC", "tot": 9e6, "exclusion_flags": [],
         "exclusion": [{"kind": "name_match", "support": "unsupported", "name": "BETA INC", "agency": "HHS", "type": "x"}],
         "signals": [{"id": "S3", "detail": "FY25 is 9x FY24"}, {"id": "R_split", "detail": "3 UEIs"}],
         "links": [{"uei": "U9", "name": "GAMMA", "via": "contact"}]}
    facts = facts_for(v, None, None, None, MAN)
    by = {f["kind"]: f for f in facts}
    assert by["name"]["weak"] and "unconfirmed" in by["name"]["text"]
    assert all(f["innocent"] for f in facts if f["level"] == "indicator")
    assert any(f["level"] == "anomaly" and "Hypergrowth" in f["text"] for f in facts)
    assert categorize(facts, "") == PATTERN
    only_name = facts_for({**v, "signals": [], "links": []}, None, None, None, MAN)
    assert categorize(only_name, "") == SINGLE


def test_cases_rank_record_backed_payments_first_and_leave_name_only_out():
    items = [
        {"v": {**EXCLUDED}, "color": ("red", ["paid after"]), "entity": ENTITY,
         "awards": {"fetched_at": "2026-10-10", "excluded_since": "2025-06-18", "paid_actions": 1, "paid_dollars": 412000.0,
                    "new_awards_after": 0, "same_day": 1, "zero_dollar": 0}},
        {"v": {"uei": "U3", "name": "BIG PATTERN CO", "tot": 9e9, "exclusion_flags": [],
               "signals": [{"id": "S3", "detail": "x"}, {"id": "R_young", "detail": "y"}]}, "color": ("red", ["Two or more"])},
        {"v": {"uei": "U4", "name": "SAME NAME LLC", "tot": 1e6, "exclusion_flags": [],
               "watchlist": [{"list": "HHS-OIG LEIE", "name": "SAME NAME", "match": "name", "detail": "NY"}]},
         "color": ("red", ["name"])},
        {"v": {"uei": "U5", "name": "CLEARED", "tot": 1e6}, "color": ("green", ["Analyst decision: Clear"]),
         "disposition": {"value": "Clear"}},
    ]
    cases, counts = build_cases(items, MAN, limit=5)
    assert [c["uei"] for c in cases] == ["U1", "U3"] and cases[0]["after_dollars"] == 412000.0
    assert counts["name_only"] == 1 and counts["cleared"] == 1 and counts["red"] == 3
    assert "$412,000" in cases[0]["headline"] and cases[1]["headline"].startswith("BIG PATTERN CO shows 2")
    assert cases[0]["timeline"][0] == ("2025-06-18", "Excluded (GSA)")


def test_the_briefing_deck_for_an_import(tmp_path, monkeypatch):
    monkeypatch.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    try:
        appmod = importlib.reload(appmod)
        c = TestClient(appmod.app)
        vendors, excl, _, _ = make_synthetic(tmp_path / "in", n=400, seed=9)
        run = appmod.store.create_run(vendors, excl, date(2026, 10, 2), synthetic=True, analyst="Ana")
        rows = appmod.store.vendors(run)["rows"]
        ex = next(r for r in rows if "EXCLUDED" in (r.get("exclusion_flags") or []))
        since = min(h["active_date"] for h in ex["exclusion"] if h["kind"] == "direct")
        since = date.fromisoformat(since).isoformat() if "-" in since else since
        d = appmod.store.root / "cases" / run / ex["uei"]
        d.mkdir(parents=True, exist_ok=True)
        act = {**ENTITY["actions"][0], "date": "2099-01-01"}
        d.joinpath("awards.json").write_text(json.dumps({"fetched_at": "2026-10-10T00:00:00", "entities": [
            {**ENTITY, "excluded_since": since, "actions": [act]}]}))
        r = c.get(f"/api/runs/{run}/exports/findings.pptx?cases=5")
        assert r.status_code == 200 and "presentationml" in r.headers["content-type"]
        deck = Presentation(io.BytesIO(r.content))
        text = "\n".join(sh.text_frame.text if sh.has_text_frame else
                         "\n".join(cell.text for row in sh.table.rows for cell in row.cells) if sh.has_table else ""
                         for s in deck.slides for sh in s.shapes)
        for want in ("Bottom line", "How to read the findings", "What existing checks miss", "Cases at a glance",
                     f"Evidence: {ex['name'][:20]}", "Suspected improper payment", "Confirmed fact",
                     "47QSWA18D0001 mod P00012", "https://www.usaspending.gov/award/CONT_IDV_1",
                     "Appendix C: Evidence index", "not legal conclusions"):
            assert want in text, want
        assert len(deck.slides) >= 14
        assert c.get("/api/runs/nope/exports/findings.pptx").status_code == 404
        assert any("Findings briefing" in e["detail"] for e in appmod.store.run_log(run))
    finally:
        monkeypatch.undo()
        importlib.reload(appmod)
