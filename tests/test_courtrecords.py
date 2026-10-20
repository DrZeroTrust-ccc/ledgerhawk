import io

from openpyxl import load_workbook

from ledgerhawk.exports.courts import build_court_records
from ledgerhawk.pipeline import context as cx
from ledgerhawk.pipeline import courtrecords as cr


def _docket(title, docket="1:24-cv-00123", suit="", cause="", parties=(), court_id="vaed", court="E.D. Va."):
    return {"source": "Federal docket", "title": title, "url": f"https://www.courtlistener.com/docket/{abs(hash(title))}/",
            "date": "2024-03-01", "where": court, "snippet": cause or suit, "tags": cx.adverse_tags(f"{title} {cause}"),
            "case": {"kind": "docket", "court": court, "court_id": court_id, "docket": docket, "filed": "2024-03-01",
                     "terminated": "", "suit": suit, "cause": cause, "parties": list(parties)}}


def _doj(title, body=""):
    return {"source": "DOJ press release", "title": title, "url": "https://www.justice.gov/opa/pr/x", "date": "2025-06-01",
            "where": "Civil Division", "snippet": body, "tags": cx.adverse_tags(f"{title} {body}")}


C = cx.clues_for("RIDGE ANALYTICS LLC", uei="SYNG6LE0B3VU", state="VA")


def test_case_types():
    assert cr.case_type(_docket("United States v. Ridge Analytics LLC", docket="1:24-cr-00077")) == cr.CRIMINAL
    assert cr.case_type(_docket("United States ex rel. Doe v. Ridge Analytics LLC", cause="31:3729 False Claims Act")) == cr.FCA
    assert cr.case_type(_docket("Ridge Analytics LLC v. United States", court_id="uscfc",
                                court="United States Court of Federal Claims")) == cr.CONTRACT_US
    assert cr.case_type(_docket("Acme Corp v. Ridge Analytics LLC", suit="190 Contract: Other")) == cr.CONTRACT
    assert cr.case_type(_docket("Smith v. Ridge Analytics LLC", suit="710 Fair Labor Standards Act")) == cr.EMPLOYMENT
    assert cr.case_type(_docket("United States v. Ridge Analytics LLC", suit="190 Contract: Other",
                                parties=["United States of America", "Ridge Analytics LLC"])) == cr.US_CIVIL
    assert cr.case_type(_doj("Virginia Contractor Agrees to Pay $2M to Resolve False Claims Act Allegations")) == cr.FCA
    assert cr.case_type(_doj("Ridge Analytics owner pleads guilty to bribery")) == cr.CRIMINAL


def test_grades_need_the_vendor_as_a_party():
    crim = _docket("United States v. Ridge Analytics LLC", docket="1:24-cr-00077")
    assert cr.grade(crim, C) == cr.GOVERNMENT and cr.role(crim, C) == "Defendant"
    # named only in the party list, not the caption
    fca = _docket("United States ex rel. Doe v. Northwind Holdings", cause="31:3729",
                  parties=["Northwind Holdings Inc", "Ridge Analytics, LLC"])
    assert cr.grade(fca, C) == cr.GOVERNMENT
    # private litigation where the vendor is the plaintiff
    suit = _docket("Ridge Analytics LLC v. Acme Corp", suit="190 Contract: Other")
    assert cr.grade(suit, C) == cr.PARTY and cr.role(suit, C) == "Plaintiff"
    # a government case about someone else that only mentions the name is not ours
    other = _docket("United States v. John Smith", docket="1:24-cr-00078", parties=["John Smith"])
    assert cr.grade(other, C) == cr.MENTION
    # a different business that shares the core name
    realty = _docket("United States v. Ridge Analytics Realty", docket="1:24-cr-00079")
    assert cr.grade(realty, cx.clues_for("RIDGE ANALYTICS LLC")) == cr.MENTION
    assert cr.grade(_doj("Ridge Analytics LLC Agrees to Pay $2M to Resolve False Claims Act Allegations"), C) == cr.GOVERNMENT
    assert cr.grade(_doj("Contractor Agrees to Pay $2M to Resolve False Claims Act Allegations", "Ridge Analytics LLC"), C) == cr.MENTION


def _entity(items_courts, items_doj=(), err=""):
    return {"name": "RIDGE ANALYTICS LLC", "uei": "SYNG6LE0B3VU", "ref": 1, "fetched_at": "2026-10-19T20:00:00+00:00",
            "clues": C, "sources": {"courts": {"items": list(items_courts), "error": err},
                                    "doj": {"items": list(items_doj), "error": ""}}}


def test_records_skip_ruled_out_and_sort_government_first():
    ruled_out = _docket("United States v. Ridge Analytics LLC", docket="1:20-cr-00001") | {"verdict": {"verdict": "not"}}
    recs, err = cr.records(_entity([
        _docket("Ridge Analytics LLC v. Acme Corp", suit="190 Contract: Other"),
        _docket("United States ex rel. Doe v. Ridge Analytics LLC", cause="31:3729 False Claims Act"),
        ruled_out,
    ]))
    assert [r["grade"] for r in recs] == [cr.GOVERNMENT, cr.PARTY] and not err
    assert recs[0]["why_it_matters"].startswith("False Claims Act")


def test_workbook():
    screen = {"meta": {"id": "S1", "matter": "GSA court review"},
              "subjects": [{"ref": 1, "entities": [{"uei": "SYNG6LE0B3VU", "tot": 2_500_000.0}]}],
              "context": {"entities": [
                  _entity([_docket("United States v. Ridge Analytics LLC", docket="1:24-cr-00077"),
                           _docket("Jones v. Smith", parties=["Jones", "Smith"])]),
                  {"name": "QUIET CO", "uei": "QUIET0000001", "fetched_at": "2026-10-19T20:00:00+00:00",
                   "sources": {"courts": {"items": [], "error": "Federal courts (CourtListener) did not answer (its rate limit was reached)"}}},
              ]}}
    wb = load_workbook(io.BytesIO(build_court_records(screen)))
    assert wb.sheetnames == ["Read Me", "Court Records", "By Vendor"]
    rows = list(wb["Court Records"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][1] == "Government case" and rows[0][5] == cr.CRIMINAL and rows[0][4] == 2_500_000
    by = list(wb["By Vendor"].iter_rows(min_row=2, values_only=True))
    assert by[0][0] == "RIDGE ANALYTICS LLC" and by[0][3] == 1
    assert "rate limit" in by[1][8]
    readme = {r[0]: r[1] for r in wb["Read Me"].iter_rows(values_only=True)}
    assert readme["Court search incomplete"].startswith("1 vendors")
    # name-only rows come in on request
    wb = load_workbook(io.BytesIO(build_court_records(screen, include_mentions=True)))
    assert len(list(wb["Court Records"].iter_rows(min_row=2))) == 2


def test_doj_releases_that_name_a_place_or_a_prison_are_not_cases():
    county = cx.clues_for("MIAMI-DADE COUNTY", state="FL")
    assert cr.grade(_doj("Five Charged with Smuggling Off the Coast of Miami-Dade County"), county) == cr.MENTION
    assert cr.grade(_doj("Miami-Dade County Man Sentenced for Fraud"), county) == cr.MENTION
    # a public body counts only with civil-enforcement or False Claims Act language
    assert cr.grade(_doj("Miami-Dade County Sentenced Officials Named in Bribery Probe"), county) == cr.MENTION
    assert cr.grade(_doj("Miami-Dade County Agrees to Pay $1M to Settle False Claims Act Allegations"), county) == cr.GOVERNMENT
    fpi = cx.clues_for("FEDERAL PRISON INDUSTRIES INC")
    assert cr.grade(_doj("Justice Department Names New Chief Executive of Federal Prison Industries Inc"), fpi) == cr.MENTION
    assert cr.grade(_doj("Federal Prison Industries Inc Agrees to Settle Kickback Claims"), fpi) == cr.GOVERNMENT
    # a firm in the title with charging language is still a government case
    assert cr.grade(_doj("Ridge Analytics LLC Indicted in Contract Fraud Scheme"), C) == cr.GOVERNMENT


def test_a_plain_phrase_name_needs_its_full_legal_name():
    fc = cx.clues_for("FEDERAL CONTRACTS LLC")
    assert cr.grade(_doj("Raytheon Agrees to Pay $950M to Resolve False Claims Act Allegations on Federal Contracts"), fc) == cr.MENTION
    assert cr.grade(_doj("Federal Contracts LLC Owner Pleads Guilty to Bribery"), fc) == cr.GOVERNMENT
