import importlib
import io
import json
import urllib.error
import urllib.parse

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ledgerhawk.pipeline import context as cx
from ledgerhawk.pipeline import osint

RSS = """<?xml version="1.0"?><rss><channel>
<item><title>Ridge Analytics owner sentenced in kickback scheme - Example Times</title><link>https://news.example/1</link>
<pubDate>Tue, 02 Sep 2025 10:00:00 GMT</pubDate><source url="https://news.example">Example Times</source></item>
<item><title>Ridge Analytics opens new office</title><link>https://news.example/2</link>
<pubDate>Mon, 05 Jan 2026 10:00:00 GMT</pubDate><source url="https://b.example">Biz Journal</source></item>
</channel></rss>"""


LEIE = (
    "LASTNAME,FIRSTNAME,MIDNAME,BUSNAME,GENERAL,SPECIALTY,UPIN,NPI,DOB,ADDRESS,CITY,STATE,ZIP,EXCLTYPE,EXCLDATE,REINDATE,WAIVERDATE,WVRSTATE\n"
    ",,,SUNRISE HOME HEALTH LLC,HOME HEALTH AGENCY,,,0,,1 MAIN ST,DAYTON,OH,45402,1128a1,20190312,00000000,00000000,\n"
)


def fake_sources(calls=None, down=()):
    """Stands in for Google News, DOJ, CourtListener, SEC EDGAR and OFAC, keyed by host."""
    def fetch(url, headers):
        host = urllib.parse.urlparse(url).netloc
        if calls is not None:
            calls.append((host, url, headers))
        if any(d in host for d in down):
            raise urllib.error.URLError("blocked")
        if "news.google.com" in host:
            return RSS.encode()
        if "justice.gov" in host:
            return json.dumps({"results": [{"title": "Virginia Contractor Agrees to Pay $2M to Resolve False Claims Act Allegations",
                                            "url": "https://www.justice.gov/opa/pr/x", "date": "1717200000",
                                            "component": [{"name": "Civil Division"}], "body": "<p>Ridge Analytics LLC...</p>"}]}).encode()
        if "courtlistener" in host:
            q = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
            if q["type"] == ["r"]:
                return json.dumps({"count": 1, "results": [{"caseName": "United States ex rel. Doe v. Ridge Analytics LLC",
                                                            "court": "E.D. Va.", "docketNumber": "1:24-cv-00123",
                                                            "dateFiled": "2024-03-01", "docket_absolute_url": "/docket/1/doe-v-ridge/",
                                                            "cause": "31:3729 False Claims Act"}]}).encode()
            return json.dumps({"count": 0, "results": []}).encode()
        if "efts.sec.gov" in host:
            return json.dumps({"hits": {"hits": [{"_id": "0001234567-25-000001:ex10.htm", "_source": {
                "display_names": ["BIGCO HOLDINGS (CIK 0000123456)"], "ciks": ["0000123456"], "form": "10-K",
                "file_date": "2025-02-20", "adsh": "0001234567-25-000001"}}]}}).encode()
        if url == cx.OFAC_URL:
            return b'36,"RIDGE ANALYTICS LLC","-0- ","SDGT",-0-\n37,"FOSTERLING, Reese","individual","RUSSIA-EO14024",-0-\n'
        if url == cx.OFAC_ALT_URL:
            return b'36,1,"aka","RIDGE DATA GROUP",-0-\n'
        if url == osint.LEIE_URL:
            return LEIE.encode()
        if "rdap.org" in host:
            return json.dumps({"events": [{"eventAction": "registration", "eventDate": "2011-04-02T00:00:00Z"}]}).encode()
        if "web.archive.org" in host:
            return json.dumps([["timestamp"], ["20110601000000"], ["20250901000000"]]).encode()
        if "bing.com" in host or "gdeltproject" in host:  # backup news feeds: unreachable unless a test provides them
            raise urllib.error.URLError("blocked")
        raise AssertionError(url)
    return fetch


@pytest.fixture(autouse=True)
def fresh_ofac(monkeypatch):
    cx._ofac_cache.update(at=0.0, rows=None)
    osint._leie.update(at=0.0, rows=None)
    cx._cl_state.update(last=0.0, blocked_until=0.0)
    monkeypatch.setenv("COURTLISTENER_GAP_S", "0")


def test_courtlistener_rate_limit_stops_further_calls(tmp_path):
    hits = []

    def fetch(url, headers):
        if "courtlistener" in url:
            hits.append(url)
            raise urllib.error.HTTPError(url, 429, "Too Many Requests", {"Retry-After": "120"}, None)
        return fake_sources([])(url, headers)

    first = cx.lookup("RIDGE ANALYTICS LLC", fetch=fetch, cache_dir=tmp_path)
    second = cx.lookup("OTHER FIRM LLC", fetch=fetch, cache_dir=tmp_path)
    assert len(hits) == 1  # one 429, then CourtListener is skipped instead of retried
    assert "rate limit" in first["sources"]["courts"]["error"]
    assert "re-check in about 2 min" in second["sources"]["courts"]["error"]
    assert not second["sources"]["doj"]["error"]  # other sources still run


def test_lookup_all_sources(tmp_path):
    calls = []
    res = cx.lookup("RIDGE ANALYTICS LLC", uei="SYNG6LE0B3VU", state="NY", fetch=fake_sources(calls), cache_dir=tmp_path)
    assert res["query"] == "Ridge Analytics" and res["errors"] == 0
    s = res["sources"]
    assert [i["date"] for i in s["news"]["items"]] == ["2025-09-02", "2026-01-05"]
    assert s["news"]["items"][0]["tags"] == ["criminal"] and s["news"]["items"][1]["tags"] == []
    assert s["news"]["items"][0]["where"] == "Example Times"
    assert s["doj"]["items"][0]["tags"] == ["civil enforcement"] and s["doj"]["items"][0]["date"] == "2024-06-01"
    assert s["courts"]["items"][0]["url"] == "https://www.courtlistener.com/docket/1/doe-v-ridge/"
    assert "civil enforcement" in s["courts"]["items"][0]["tags"]  # qui tam / false claims in the cause
    assert s["sec"]["items"][0]["url"] == "https://www.sec.gov/Archives/edgar/data/123456/000123456725000001/ex10.htm"
    assert [i["title"] for i in s["ofac"]["items"]] == ["RIDGE ANALYTICS LLC"]  # alias row is a different name
    assert res["adverse"] == 4 and res["count"] == 6
    assert any("opencorporates" in m["url"] and "us_ny" in m["url"] for m in res["manual"])
    assert all(c[0] != "news.google.com" or '%22Ridge+Analytics%22' in c[1] for c in calls)  # exact-phrase search
    assert not any("Authorization" in c[2] for c in calls if "courtlistener" in c[0])
    assert (tmp_path / "ofac.json").exists()  # sanctions list cached for the day


def test_person_and_failures(tmp_path, monkeypatch):
    monkeypatch.setenv("COURTLISTENER_TOKEN", "tok")
    calls = []
    res = cx.lookup("reese fosterling", person=True, fetch=fake_sources(calls, down=("efts.sec.gov",)), cache_dir=tmp_path)
    assert res["query"] == "Reese Fosterling"
    assert res["sources"]["ofac"]["items"][0]["title"] == "FOSTERLING, Reese"  # 'LAST, First' matched in any order
    assert res["sources"]["sec"]["error"].startswith("SEC filings did not answer") and res["errors"] == 1
    assert all(c[2].get("Authorization") == "Token tok" for c in calls if "courtlistener" in c[0])
    assert not any("opencorporates" in m["url"] for m in res["manual"])
    assert cx.adverse_tags("The agency issued guidance") == [] and cx.adverse_tags("Vendor sues GSA") == ["litigation"]


def _client(tmp_path):
    mp = pytest.MonkeyPatch()
    mp.setenv("LEDGERHAWK_DATA_DIR", str(tmp_path / "data"))
    mp.setenv("LEDGERHAWK_WEB_DIST", str(tmp_path / "no-web"))
    import ledgerhawk.api.app as appmod
    appmod = importlib.reload(appmod)
    return mp, appmod, TestClient(appmod.app)


def test_context_api_and_subject_exports(tmp_path):
    from docx import Document
    from ledgerhawk.pipeline.synthetic import make_synthetic
    d = tmp_path / "syn"
    d.mkdir()
    vendors, excl, sam, p = make_synthetic(d, n=1500, seed=5)
    mp, appmod, client = _client(tmp_path)
    try:
        appmod.store.context_fetch = fake_sources(down=tuple(["news.google", "justice.gov", "courtlistener", "efts.sec.gov",
                                                               "ofac", "oig.hhs.gov"]))
        assert client.post("/api/context", data={"analyst": "A", "name": "Ridge Analytics LLC"}).status_code == 502
        appmod.store.context_fetch = fake_sources()
        assert client.post("/api/context", data={"analyst": "", "name": "X Y"}).status_code == 400
        assert client.get("/api/context", params={"uei": "SYNG6LE0B3VU"}).json()["context"] is None
        r = client.post("/api/context", data={"analyst": "A", "name": "Ridge Analytics LLC", "uei": "SYNG6LE0B3VU"})
        assert r.status_code == 200, r.text
        got = client.get("/api/context", params={"uei": "SYNG6LE0B3VU"}).json()["context"]
        assert got["fetched_by"] == "A" and got["adverse"] == 4

        ids = {}
        for kind, path, dt in (("sam", sam, "2026-09-06"), ("exclusions", excl, "2026-10-02")):
            with open(path, "rb") as f:
                ids[kind] = client.post("/api/sources", files={"file": f}, data={"kind": kind, "as_of": dt, "analyst": "T"}).json()["id"]
        sid = client.post("/api/subject-screens", data={"subjects_text": p["ex_affiliate"], "people_text": "Reese Fosterling, NY",
                                                        "analyst": "T", "sam_source": ids["sam"],
                                                        "exclusions_source": ids["exclusions"]}).json()["id"]
        r = client.post(f"/api/subject-screens/{sid}/context", data={"analyst": "A"})
        assert r.status_code == 200, r.text
        ents = client.get(f"/api/subject-screens/{sid}").json()["context"]["entities"]
        assert [e.get("ref") for e in ents] == [1, None] and ents[1]["person_ref"] == 1 and ents[1]["person"]
        wb = load_workbook(io.BytesIO(client.get(f"/api/subject-screens/{sid}/subject-screen.xlsx").content))
        ws = wb["Outside Context"]
        assert ws["B6"].value == "Subject 1" and ws["E6"].value and ws["D6"].value is not None  # every row says how it matched
        doc = Document(io.BytesIO(client.get(f"/api/subject-screens/{sid}/subject-screen.docx").content))
        text = "\n".join(par.text for par in doc.paragraphs)
        # The fake news is about Ridge Analytics, not this subject: a name-only hit stays out of the written report.
        assert text.count("Outside context") == 2 and "kickback scheme" not in text and "name-only (not listed)" in text
        assert "FOSTERLING, Reese" in text  # the same-name sanctions entry is reported, labeled unverified
        assert "[Possible match 50/100, not verified] OFAC SDN list" in text
        assert any(a["action"] == "context_lookup" for a in client.get("/api/audit").json())
    finally:
        mp.undo()


def test_parallel_lookups_share_one_ofac_download(tmp_path):
    calls = []
    found = cx.lookup_many([{"name": f"Firm Number {n} LLC"} for n in range(6)], fetch=fake_sources(calls), cache_dir=tmp_path)
    assert all(f["errors"] == 0 for f in found)
    assert sum(1 for c in calls if c[1] == cx.OFAC_URL) == 1


def test_match_strength():
    c = cx.clues_for("RIDGE ANALYTICS LLC", uei="SYNG6LE0B3VU", state="VA", city="Reston", cage="7ABC1",
                     people=["Dana Whitfield"], related=["Ridge Data Partners LLC"])
    def conf(title, snippet="", where=""):
        return cx.score({"title": title, "snippet": snippet, "where": where, "match": "name"}, c)["confidence"]
    assert conf("Ridge Analytics LLC wins Army task order", where="Reston, Virginia") == "strong"  # name and city
    assert conf("Ridge Analytics CEO Dana Whitfield indicted") == "strong"  # name and an officer
    assert conf("Contract to CAGE 7ABC1 protested") == "strong"  # an identifier alone is decisive
    assert conf("Ridge Analytics opens Denver office") == "possible"  # distinctive name, nothing else
    assert conf("Ridge Analytics sued", where="E.D. Va.") == "possible"  # the state alone is a weak tie
    assert conf("Local ridge hiking analytics app launches") == "weak"  # words apart, no corroboration
    other = cx.score({"title": "Smith v. Ridge Analytics Realty LLC", "snippet": "", "where": "", "match": "name"}, c)
    assert other["confidence"] == "weak" and other["score"] == 10
    assert other["why"][0] == "names a different business (Ridge Analytics Realty), so the score is capped at 10"
    assert conf("Ridge Analytics Inc. hires CFO") == "possible"  # a legal suffix is not another business
    assert conf("10-K filed by BIGCO HOLDINGS") == "weak"  # name not shown at all
    g = cx.clues_for("GLOBAL SOLUTIONS GROUP LLC")
    assert g["generic"]
    hit = cx.score({"title": "Global Solutions Group fined", "snippet": "", "where": "", "match": "name"}, g)
    assert hit["confidence"] == "weak" and "common business name" in hit["why"]
    p = cx.clues_for("Reese Fosterling", state="NY", person=True, related=["Ridge Analytics LLC"])
    pc = lambda t: cx.score({"title": t, "snippet": "", "where": "", "match": "name"}, p)["confidence"]  # noqa: E731
    assert pc("Reese Fosterling named in suit") == "weak"  # a person's name alone proves little
    assert pc("New York man Reese Fosterling charged") == "possible"
    assert pc("Ridge Analytics founder Reese Fosterling charged") == "strong"


def test_ofac_type_must_match(tmp_path):
    # An individual named like a company, or a company named like a person, is not a match.
    rows = b'40,"HARBOR LANE","individual","SDGT",-0-\n41,"HARBOR LANE","-0- ","SDGT",-0-\n'
    fetch = lambda url, h: rows if url == cx.OFAC_URL else b""  # noqa: E731
    assert [i["title"] for i in cx.ofac("Harbor Lane LLC", fetch, tmp_path)] == ["HARBOR LANE"]
    assert len(cx.ofac("Harbor Lane", fetch, tmp_path, person=True)) == 1
    assert cx.ofac("Harbor Lane LLC", fetch, tmp_path)[0]["where"].startswith("entity")


def test_verdicts(tmp_path):
    mp, appmod, client = _client(tmp_path)
    try:
        appmod.store.context_fetch = fake_sources()
        q = {"uei": "SYNG6LE0B3VU", "name": "Ridge Analytics LLC"}
        got = client.post("/api/context", data={"analyst": "A", **q, "state": "VA", "city": "Reston",
                                                "people": "Dana Whitfield\n"}).json()
        assert got["clues"]["city"] == "Reston" and got["clues"]["people"] == ["Dana Whitfield"]
        items = {i["title"]: i for s in got["sources"].values() for i in s["items"]}
        doj = next(i for t, i in items.items() if t.startswith("Virginia Contractor"))
        assert doj["confidence"] == "possible" and "mentions state Virginia (+10)" in doj["why"] and doj["score"] == 60
        news = items["Ridge Analytics opens new office"]
        r = client.post("/api/context/verdict", data={"analyst": "B", **q, "item": news["id"], "verdict": "not"})
        assert r.status_code == 400  # ruling a hit out needs a reason
        r = client.post("/api/context/verdict", data={"analyst": "B", **q, "item": news["id"], "verdict": "not",
                                                      "note": "Different company in Ohio"})
        assert r.status_code == 200, r.text
        r = client.post("/api/context/verdict", data={"analyst": "B", **q, "item": doj["id"], "verdict": "same"})
        t = r.json()["tally"]
        assert t["dismissed"] == 1 and t["confirmed"] == 1
        # A refresh keeps both calls.
        client.post("/api/context", data={"analyst": "A", **q})
        again = client.get("/api/context", params={"uei": q["uei"]}).json()["context"]
        assert again["tally"]["dismissed"] == 1 and again["tally"]["confirmed"] == 1
        ranked = cx.ranked(again)
        assert ranked[0]["id"] == doj["id"] and ranked[-1]["id"] == news["id"]
        from ledgerhawk.exports.subjects import context_items, context_line
        assert news["id"] not in [i["id"] for i in context_items(again)]
        assert context_line(ranked[0]).startswith("[Confirmed by B]")
        assert client.post("/api/context/verdict", data={"analyst": "B", **q, "item": "nope", "verdict": "same"}).status_code == 404
        log = [a for a in client.get("/api/audit").json() if a["action"] == "context_verdict"]
        assert len(log) == 2 and any("Different company in Ohio" in a["detail"] for a in log)
    finally:
        mp.undo()


BING = b"""<?xml version="1.0"?><rss xmlns:News="https://www.bing.com/news/search?q=x&amp;format=rss"><channel>
<item><title>Ridge Analytics settles false claims case</title><link>https://bing.example/1</link>
<description>Ridge Analytics LLC of Reston agreed...</description><pubDate>Tue, 02 Sep 2025 10:00:00 GMT</pubDate>
<News:Source>Reston Now</News:Source></item></channel></rss>"""


def test_news_falls_back_when_google_refuses(tmp_path):
    base = fake_sources(down=("news.google.com",))

    def fetch(url, headers):
        if "bing.com" in url:
            return BING
        return base(url, headers)
    items = cx.news("Ridge Analytics", fetch)
    assert [(i["title"], i["where"], i["date"]) for i in items] == [("Ridge Analytics settles false claims case", "Reston Now", "2025-09-02")]
    assert "Reston" in items[0]["snippet"]

    def gdelt_only(url, headers):
        if "gdeltproject" in url:
            return json.dumps({"articles": [{"title": "Ridge Analytics wins award", "url": "https://g.example/1",
                                             "seendate": "20260105T101500Z", "domain": "g.example"}]}).encode()
        raise urllib.error.HTTPError(url, 503, "busy", {}, None)
    assert [(i["date"], i["where"]) for i in cx.news("Ridge Analytics", gdelt_only)] == [("2026-01-05", "g.example")]

    def none(url, headers):
        raise urllib.error.HTTPError(url, 503, "busy", {}, None)
    res = cx.lookup("Ridge Analytics LLC", fetch=lambda u, h: none(u, h) if "efts" in u or "news" in u or "bing" in u or "gdelt" in u
                    else fake_sources()(u, h), cache_dir=tmp_path)
    assert res["sources"]["news"]["error"] == "News did not answer (Google News 503; Bing News 503; GDELT 503)"


def test_fetch_retries_once_on_server_error(monkeypatch):
    calls = []

    class Resp(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    def urlopen(req, timeout):
        calls.append(req.full_url)
        if len(calls) == 1:
            raise urllib.error.HTTPError(req.full_url, 500, "oops", {}, None)
        return Resp(b"ok")
    monkeypatch.setattr(cx.urllib.request, "urlopen", urlopen)
    monkeypatch.setattr(cx.time, "sleep", lambda s: None)
    assert cx._fetch("https://efts.sec.gov/x", {}) == b"ok" and len(calls) == 2
    calls.clear()

    def forbidden(req, timeout):
        calls.append(1)
        raise urllib.error.HTTPError(req.full_url, 403, "no", {}, None)
    monkeypatch.setattr(cx.urllib.request, "urlopen", forbidden)
    with pytest.raises(urllib.error.HTTPError):
        cx._fetch("https://efts.sec.gov/x", {})
    assert calls == [1]  # a refusal is not retried


def test_brave_search(tmp_path, monkeypatch):
    monkeypatch.setenv("BRAVE_API_KEY", "k")
    monkeypatch.setenv("BRAVE_QPS", "1000")
    calls = []
    base = fake_sources()

    def fetch(url, headers):
        if "search.brave.com" not in url:
            return base(url, headers)
        calls.append((url, headers))
        q = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)["q"][0]
        row = {"title": f"<strong>Ridge Analytics</strong> result for {q[:20]}", "url": f"https://x.example/{len(calls)}",
               "description": "Ridge Analytics LLC of Reston, Virginia, agreed to settle", "page_age": "2025-09-02T10:00:00",
               "meta_url": {"hostname": "x.example"}}
        dup = {"title": "dup", "url": "https://x.example/1", "description": ""}
        return json.dumps({"web": {"results": [row, dup]}} if "/web/" in url else {"results": [row]}).encode()
    c = cx.clues_for("RIDGE ANALYTICS LLC", state="VA", city="Reston", people=["Dana Whitfield"])
    res = cx.lookup("RIDGE ANALYTICS LLC", clues=c, fetch=fetch, cache_dir=tmp_path)
    items = res["sources"]["brave"]["items"]
    assert res["web_search"] and len(calls) == 5  # three web queries, news for the first two
    assert all(h["X-Subscription-Token"] == "k" for _, h in calls)
    assert len(items) == 5 and items[0]["title"].startswith("Ridge Analytics result")  # tags stripped, duplicate dropped
    assert items[0]["confidence"] == "strong" and items[0]["date"] == "2025-09-02"  # Reston in the description
    assert items[0]["query"] == '"Ridge Analytics" Reston'
    assert any('"Dana Whitfield" "Ridge Analytics"' in urllib.parse.unquote_plus(u) for u, _ in calls)
    monkeypatch.delenv("BRAVE_API_KEY")
    off = cx.lookup("RIDGE ANALYTICS LLC", fetch=fetch, cache_dir=tmp_path)
    assert "brave" not in off["sources"] and off["web_search"] is False


def test_scores_and_junk_controls(tmp_path):
    c = cx.clues_for("RIDGE ANALYTICS LLC", uei="SYNG6LE0B3VU", state="VA", city="Reston", people=["Dana Whitfield"])
    def sc(title, url="https://news.example/x", snippet=""):
        return cx.score({"title": title, "snippet": snippet, "where": "", "url": url, "match": "name"}, c)
    assert sc("Ridge Analytics opens office")["score"] == 40
    assert sc("Ridge Analytics opens office", url="https://www.bizapedia.com/va/ridge")["score"] == 25  # directory site
    assert sc("Ridge Analytics LLC sued in Reston", url="https://www.justice.gov/x")["score"] == 80
    assert sc("Ridge Analytics CEO Dana Whitfield, Reston", snippet="UEI SYNG6LE0B3VU")["score"] == 99  # capped
    assert sc("Ridge Analytics opens office")["why"] == ["name without legal suffix (+40)"]

    mp, appmod, client = _client(tmp_path)
    try:
        appmod.store.context_fetch = fake_sources()
        q = {"uei": "SYNG6LE0B3VU", "name": "Ridge Analytics LLC"}
        got = client.post("/api/context", data={"analyst": "A", **q}).json()
        news = [i for i in got["sources"]["news"]["items"]]
        # Bulk rule-out of several hits at once.
        r = client.post("/api/context/verdict", data={"analyst": "B", **q, "item": ",".join(i["id"] for i in news),
                                                      "verdict": "not", "note": "Name only, ruled out in bulk"})
        assert r.status_code == 200, r.text
        assert r.json()["tally"]["dismissed"] == len(news)
        # Muting a site rules out its hits everywhere, unless someone decided one individually.
        assert client.post("/api/context/muted-sites", data={"analyst": "B", "host": "not a host"}).status_code == 400
        r = client.post("/api/context/muted-sites", data={"analyst": "B", "host": "www.CourtListener.com", "note": "noise"})
        assert r.json()["sites"]["courtlistener.com"]["by"] == "B"
        cxt = client.get("/api/context", params={"uei": q["uei"]}).json()["context"]
        court = cxt["sources"]["courts"]["items"][0]
        assert court["verdict"]["muted"] and court["verdict"]["note"] == "Site courtlistener.com muted: noise"
        client.post("/api/context/muted-sites", data={"analyst": "B", "host": "courtlistener.com", "mute": "false"})
        cxt = client.get("/api/context", params={"uei": q["uei"]}).json()["context"]
        assert cxt["sources"]["courts"]["items"][0]["verdict"] is None
        assert client.get("/api/context/muted-sites").json()["sites"] == {}
        log = [a["action"] for a in client.get("/api/audit").json()]
        assert log.count("context_mute_site") == 2
    finally:
        mp.undo()
