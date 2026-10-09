"""The FAR referral memo: each provision to review, element by element, with the analyst's notes and the import's
provenance, for the suspension and debarment official or the OIG."""
import io
import json
from datetime import datetime, timezone

from docx import Document

from ledgerhawk.exports.word import FAR_NOT_A_FINDING, build_far_memo_docx
from ledgerhawk.pipeline.far import apply_decisions, evaluate
from ledgerhawk.pipeline.integrity import _parse_date
from tests.test_far import EXCL, _act, _v

SUMMARY = {"meta": {"id": "run-7", "label": "FY25 pilot", "created_at": "2026-10-03T12:00:00", "data_class": "synthetic"},
           "manifest": {"rule_set_version": "2026.10-pilot", "rule_set_fingerprint": "fp-abc123", "input_file": "vendors.xlsx",
                        "input_sha256": "deadbeef", "sam_extract_date": "2026-09-06", "exclusions_extract_date": "2026-10-02"}}


def _text(body: bytes) -> str:
    doc = Document(io.BytesIO(body))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_memo_lists_provisions_elements_notes_and_provenance():
    v = _v(exclusion_flags=["EXCLUDED"], exclusion=[EXCL], fy24=1e5, fy25=2e5, tot=3e5)
    provs = evaluate(v, {"actions": [_act("2025-07-01", "new", 2e4), _act("2025-08-01", "option", 3e4)], "awards": []})
    at = "2026-10-08T15:00:00"
    far = apply_decisions(provs, {
        "9.405|determination": {"state": "confirmed", "note": "No D&F in the contract file", "analyst": "Pat", "at": at},
        "52.209-5|certification": {"state": "not_applicable", "note": "Offer predates SAM reps", "analyst": "Pat", "at": at},
        "9.405-1|*": {"state": "not_applicable", "note": "Option was exercised by another agency", "analyst": "Lee", "at": at},
    })
    body = build_far_memo_docx(v, far, SUMMARY, matter="M-12", privileged=True, usaspending_at="2026-10-05T09:00:00",
                               generated_at=datetime(2026, 10, 9, tzinfo=timezone.utc))
    doc = Document(io.BytesIO(body))
    text = _text(body)
    assert FAR_NOT_A_FINDING in text and "Matter: M-12" in text and "SYNTHETIC DATA" in text
    assert "Privileged and Confidential" in doc.sections[0].header.paragraphs[0].text
    assert doc.core_properties.author == "LedgerHawk"
    # the provisions to review, with whom they bind and where they route
    heads = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert "FAR 9.405(a): New contract or order to an excluded contractor" in heads
    assert "FAR 52.209-5: Certification regarding responsibility matters" in heads
    assert not any(h.startswith("FAR 9.405-1(b)") for h in heads)  # set aside as a whole: not a provision to review
    assert "Awarding agency" in text and "GSA suspension and debarment official" in text and "OIG (referral after analyst review)" in text
    # every element, with its state, detail and source
    rows = {r.cells[0].text: [c.text for c in r.cells] for t in doc.tables for r in t.rows if len(r.cells) == 4}
    assert rows["A new contract or order dated after the exclusion date"][1:] == [
        "Data shows", rows["A new contract or order dated after the exclusion date"][2], "USAspending"]
    assert rows["A new contract or order dated after the exclusion date"][2].startswith("1 new contract or order")
    assert rows["No written compelling-reason determination by the agency head (FAR 9.405(a))"][1] == "Analyst confirmed"
    assert rows["What the firm certified about its debarment or suspension, in the offer or SAM reps and certs"][1] == "Not applicable"
    # the analyst's notes, and the set-aside provision with its reason
    assert "Analyst confirmed: No D&F in the contract file (Pat, 2026-10-08)" in text
    assert "Not applicable: Offer predates SAM reps (Pat, 2026-10-08)" in text
    assert "Set aside by an analyst" in heads and "Option was exercised by another agency (Lee, 2026-10-08)" in text
    # provenance
    for s in ("FY25 pilot (run-7)", "fp-abc123", "2026.10-pilot", "SHA-256 deadbeef", "2026-09-06", "2026-10-02", "2026-10-05"):
        assert s in text, s


def test_far_memo_through_the_api(sam_ctx_export):
    import ledgerhawk.api.app as appmod
    client, run_id = sam_ctx_export
    st = appmod.store
    rows = st.vendors(run_id)["rows"]
    target = next(r for r in rows if "EXCLUDED" in r["exclusion_flags"]
                  and any(h["kind"] == "direct" and _parse_date(h["active_date"]) for h in r["exclusion"]))
    u = target["uei"]
    since = min(_parse_date(h["active_date"]) for h in target["exclusion"] if h["kind"] == "direct" and _parse_date(h["active_date"]))
    d = st._case_dir(run_id, u)
    d.mkdir(parents=True, exist_ok=True)
    after = since.replace(year=since.year + 1).isoformat()
    (d / "awards.json").write_text(json.dumps({"fetched_at": "2026-10-07T10:00:00",
                                               "entities": [{"actions": [_act(after, "new", 5e4)], "awards": []}]}))
    client.post(f"/api/runs/{run_id}/vendors/{u}/far", json={"provision": "9.405", "element": "determination", "state": "confirmed",
                                                            "note": "Contract file has no D&F", "analyst": "A"})
    r = client.get(f"/api/runs/{run_id}/vendors/{u}/far-memo.docx", params={"matter": " M-3 "})
    assert r.status_code == 200 and "FAR memo" in r.headers["content-disposition"]
    text = _text(r.content)
    assert "FAR 9.405(a)" in text and "Contract file has no D&F" in text and "Matter: M-3" in text
    summary = st.summary(run_id)
    assert summary["meta"]["id"] in text and summary["manifest"]["rule_set_fingerprint"] in text and "2026-10-07" in text
    log = client.get(f"/api/runs/{run_id}/vendors/{u}").json()["history"]
    assert any(h["action"] == "export" and "FAR referral memo" in h["detail"] for h in log)
    # nothing to refer: no provisions at all, or every one set aside
    plain = next(x for x in rows if not appmod._far_lens(run_id)(x))
    assert client.get(f"/api/runs/{run_id}/vendors/{plain['uei']}/far-memo.docx").status_code == 404
    assert client.get(f"/api/runs/{run_id}/vendors/NOPE/far-memo.docx").status_code == 404
    for p in appmod._far_lens(run_id)(st.vendors(run_id)["by_uei"][u]):
        client.post(f"/api/runs/{run_id}/vendors/{u}/far", json={"provision": p["id"], "element": "*", "state": "not_applicable",
                                                                "note": "Different entity", "analyst": "A"})
    assert client.get(f"/api/runs/{run_id}/vendors/{u}/far-memo.docx").status_code == 404


from tests.test_analysis_export import sam_ctx_export  # noqa: E402,F401  (the same synthetic import)
